"""
Historial de EPP: una fila por trabajador y EPP, una columna por entrega, con
los días que duró cada una. Lo ven el SuperAdmin, el Coordinador y el
encargado de almacén, en la app y en Excel.
"""
import io
from datetime import date, datetime
from decimal import Decimal

from django.core.management import call_command
from django.utils import timezone
from openpyxl import load_workbook

from ..historial_epp import armar_historial, texto_de_entrega
from ..models import (EPP, DetallePedidoEPP, PedidoEPP, Rol, Usuario)
from .base import BaseAPITestCase


def entregar(usuario, epp, dia, tipo='nuevo', estado='aprobado', cantidad=1):
    pedido = PedidoEPP.objects.create(usuario=usuario, tipo=tipo, estado=estado)
    PedidoEPP.objects.filter(pk=pedido.pk).update(
        fecha_aprobacion=timezone.make_aware(datetime.combine(dia, datetime.min.time())))
    DetallePedidoEPP.objects.create(pedido=pedido, epp=epp,
                                    cantidad_solicitada=cantidad,
                                    cantidad_aprobada=cantidad)


class HistorialEPPTests(BaseAPITestCase):

    def setUp(self):
        super().setUp()
        call_command('configurar_epp', stdout=io.StringIO())
        self.pedro = Usuario.objects.create(
            nombre='Pedro Operario', email='pedro@x.com', clave='x', dni='4567',
            empresa=self.empresa, rol=Rol.objects.get(id_rol=Rol.OPERARIO))
        self.coordinador = Usuario.objects.create(
            nombre='Coordinadora', email='coord@x.com', clave='x',
            empresa=self.empresa,
            rol=Rol.objects.create(id_rol=Rol.COORDINADOR, descripcion='Coordinador'))
        b42 = EPP.objects.get(codigo='BOTA-DIEL-42')
        b43 = EPP.objects.get(codigo='BOTA-DIEL-43')
        casco = EPP.objects.get(codigo='CASCO-AZUL')
        entregar(self.pedro, b42, date(2026, 6, 1))
        entregar(self.pedro, b43, date(2026, 8, 15), tipo='cambio')
        entregar(self.pedro, casco, date(2026, 7, 1))
        # Rechazado: no es una entrega.
        entregar(self.pedro, casco, date(2026, 9, 1), estado='rechazado')

    def test_una_fila_por_trabajador_y_epp_con_lo_que_duro(self):
        filas = armar_historial(hoy=date(2026, 10, 5))
        self.assertEqual([f['epp'] for f in filas],
                         ['BOTAS DIELECTRICAS', 'CASCO AZUL'])
        botas = filas[0]
        self.assertEqual(botas['trabajador'], 'Pedro Operario')
        self.assertEqual([(e['fecha'], e['talla'], e['dias'], e['en_uso'])
                          for e in botas['entregas']], [
            (date(2026, 6, 1), '42', 75, False),
            (date(2026, 8, 15), '43', 51, True),
        ])
        self.assertEqual(len(filas[1]['entregas']), 1)

    def test_texto_de_la_celda(self):
        e = {'fecha': date(2026, 8, 15), 'talla': '43', 'cantidad': Decimal('1'),
             'dias': 51, 'en_uso': True}
        self.assertEqual(texto_de_entrega(e), '15/08/2026 · T43 · 51 días (en uso)')
        e.update(talla='', cantidad=Decimal('2'), dias=1, en_uso=False)
        self.assertEqual(texto_de_entrega(e), '15/08/2026 · x2 · 1 día')

    def test_lo_ven_el_coordinador_y_el_almacen(self):
        for quien in (self.coordinador, self.encargado):
            self.auth(quien)
            r = self.client.get('/api/pedidos-epp/historial/')
            self.assertEqual(r.status_code, 200, r.content)
            self.assertEqual(r.json()['columnas'], 2)
            self.assertEqual(len(r.json()['filas']), 2)
        for quien in (self.capataz, self.pedro):
            self.auth(quien)
            self.assertEqual(
                self.client.get('/api/pedidos-epp/historial/').status_code, 403,
                quien.email)

    def test_el_excel(self):
        self.auth(self.coordinador)
        r = self.client.get('/api/pedidos-epp/historial_excel/')
        self.assertEqual(r.status_code, 200)
        self.assertIn('historial_epp_', r['Content-Disposition'])
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual([c.value for c in ws[3]][:6],
                         ['Trabajador', 'DNI', 'EPP', 'Entregas',
                          'Entrega 1', 'Entrega 2'])
        self.assertEqual([c.value for c in ws[4]][:5],
                         ['Pedro Operario', '4567', 'BOTAS DIELECTRICAS', 2,
                          '01/06/2026 · T42 · 75 días'])
