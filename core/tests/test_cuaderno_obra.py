"""
Cuaderno de obra y Excel de liquidación de una SST.

El cuaderno se redacta solo con lo liquidado y lo dibujado en el plano, así que
lo que importa es que cada renglón aparezca cuando corresponde y no cuando no.
El Excel es la plantilla de Tecsur: las cantidades tienen que caer en su celda.
"""
import io
from datetime import date, time
from decimal import Decimal

from django.test import SimpleTestCase
from openpyxl import Workbook, load_workbook

from ..cuaderno_obra import (
    PARTIDAS_RETIRO_POSTE, Item, Liquidado, lineas_del_cuaderno,
)
from ..excel_liquidacion import (
    MO_ENCABEZADO, MO_POSTES, PLANTILLA, _columna_con, _columnas_de_postes,
    _rotular_postes,
)
from ..models import (
    Actividad, ActividadTipoTrabajo, ConsumoMaterialSuministro, CuadernoObra,
    LiquidacionPartida,
    LiquidacionSuministro, ManoDeObra, Material, PlanoSST, Recupero, SST,
    Suministro, SSTSuministro, SuministroRecupero, TipoTrabajo,
)
from .base import BaseAPITestCase


def item(descripcion, cantidad=1, codigo='X'):
    return Item(codigo, descripcion, Decimal(str(cantidad)))


def partida(codigo, cantidad=1):
    return Item(codigo, f'PARTIDA {codigo}', Decimal(str(cantidad)))


def cable(estado, metros=10, descripcion='CAAIS 3x16', pendiente=None):
    e = {'tipo': 'cable', 'estado': estado, 'metros': metros,
         'descripcion': descripcion}
    if pendiente is not None:
        e['pendiente'] = pendiente
    return e


def vereda(largo, ancho):
    return {'tipo': 'vereda', 'largo': largo, 'ancho': ancho}


def poste_liquidado(numero, materiales=(), partidas=(), recuperos=(), tipos=()):
    """Un poste de la SST con lo suyo. El plano no va aquí: es de la SST."""
    return Liquidado(
        materiales=list(materiales), partidas=list(partidas),
        recuperos=list(recuperos), postes=[numero], tipos=list(tipos))


class LineasDelCuadernoTests(BaseAPITestCase):
    """El cuerpo del cuaderno: cada renglón cuando corresponde, y no cuando no.

    Las reglas las dio el usuario el 2026-09-21, una por una; los textos son
    los suyos.
    """

    def test_empieza_con_la_frase_del_formato(self):
        self.assertEqual(
            lineas_del_cuaderno(Liquidado())[0],
            'Por la presente se informa que la SST se ejecutó según lo detallado:')

    def test_sin_nada_liquidado_solo_queda_la_frase(self):
        self.assertEqual(len(lineas_del_cuaderno(Liquidado())), 1)

    # ── El poste ────────────────────────────────────────────────────────────

    def test_el_poste_se_nombra_como_en_obra(self):
        d = Liquidado(materiales=[item('POSTE DE POLIESTER 7,5 / 150',
                                       codigo='5331596')])
        self.assertIn('Se instaló PRFV 7/150', lineas_del_cuaderno(d))

    def test_el_poste_de_nueve_metros(self):
        d = Liquidado(materiales=[item('POSTE POLIESTER 9 / 200',
                                       codigo='5331616')])
        self.assertIn('Se instaló PRFV 9/200', lineas_del_cuaderno(d))

    def test_cimentado_si_se_liquido_la_cimentacion(self):
        d = Liquidado(materiales=[item('POSTE', codigo='5331616')],
                      partidas=[partida('*094918')])
        self.assertIn('Se instaló PRFV 9/200 Cimentado', lineas_del_cuaderno(d))

    def test_cimentado_tambien_con_la_instalacion_con_cimentacion(self):
        d = Liquidado(materiales=[item('POSTE', codigo='5331616')],
                      partidas=[partida('*090482')])
        self.assertIn('Se instaló PRFV 9/200 Cimentado', lineas_del_cuaderno(d))

    def test_sin_cimentacion_no_lo_dice(self):
        d = Liquidado(materiales=[item('POSTE', codigo='5331616')])
        lineas = lineas_del_cuaderno(d)
        self.assertIn('Se instaló PRFV 9/200', lineas)
        self.assertFalse(any('Cimentado' in l for l in lineas))

    def test_el_codigo_del_plano_va_al_final(self):
        d = Liquidado(
            materiales=[item('POSTE', codigo='5331616')],
            partidas=[partida('*094918')],
            elementos_plano=[{'assetId': 'poste_nuevo', 'codigo': '123456789'}])
        self.assertIn('Se instaló PRFV 9/200 Cimentado con código 123456789',
                      lineas_del_cuaderno(d))

    def test_la_abrazadera_de_poste_no_es_un_poste(self):
        d = Liquidado(materiales=[item('ABRAZADERA POSTE 150MM', codigo='6941170')])
        self.assertFalse(any('Se instaló' in l for l in lineas_del_cuaderno(d)))

    # ── Alumbrado instalado ─────────────────────────────────────────────────

    def test_luminaria_instalada(self):
        d = Liquidado(partidas=[partida('*091320')])
        self.assertIn('Se instaló luminaria', lineas_del_cuaderno(d))

    def test_pastoral_instalado(self):
        d = Liquidado(partidas=[partida('*091346')])
        self.assertIn('Se instaló pastoral', lineas_del_cuaderno(d))

    def test_mas_de_uno_lleva_su_cantidad(self):
        d = Liquidado(partidas=[partida('*091320', 2)])
        self.assertIn('Se instaló luminaria (2)', lineas_del_cuaderno(d))

    def test_instalar_y_trasladar_son_renglones_distintos(self):
        d = Liquidado(partidas=[partida('*091320'), partida('*091322')])
        lineas = lineas_del_cuaderno(d)
        self.assertIn('Se instaló luminaria', lineas)
        self.assertIn('Se trasladó luminaria', lineas)
        self.assertLess(lineas.index('Se instaló luminaria'),
                        lineas.index('Se trasladó luminaria'))

    # ── Traslados de alumbrado ──────────────────────────────────────────────

    def test_pastoral_y_luminaria_van_en_un_solo_renglon(self):
        d = Liquidado(partidas=[partida('*091356'), partida('*091322')])
        self.assertIn('Se trasladó pastoral + luminaria existente',
                      lineas_del_cuaderno(d))

    def test_solo_el_pastoral(self):
        d = Liquidado(partidas=[partida('*091356')])
        self.assertIn('Se trasladó pastoral', lineas_del_cuaderno(d))

    def test_solo_la_luminaria(self):
        d = Liquidado(partidas=[partida('*091322')])
        self.assertIn('Se trasladó luminaria', lineas_del_cuaderno(d))

    def test_sin_traslado_no_se_menciona(self):
        self.assertFalse(
            any('trasladó' in l for l in lineas_del_cuaderno(Liquidado())))

    # ── La subida al poste ──────────────────────────────────────────────────

    def test_la_subida_por_su_partida(self):
        d = Liquidado(partidas=[partida('*091240')])
        self.assertIn('Se instaló Subida AP N2XY 2-1x6', lineas_del_cuaderno(d))

    def test_la_subida_por_el_metrado_del_cable(self):
        d = Liquidado(materiales=[item('CABLE N2XY 2 - 1 X 6MM2', 8,
                                       codigo='5031165')])
        self.assertIn('Se instaló Subida AP N2XY 2-1x6 (8 metros)',
                      lineas_del_cuaderno(d))

    def test_un_tendido_largo_no_es_una_subida(self):
        d = Liquidado(materiales=[item('CABLE N2XY', 40, codigo='5031165')])
        self.assertFalse(any('Subida' in l for l in lineas_del_cuaderno(d)))

    def test_un_metro_suelto_tampoco(self):
        d = Liquidado(materiales=[item('CABLE N2XY', 1, codigo='5031165')])
        self.assertFalse(any('Subida' in l for l in lineas_del_cuaderno(d)))

    def test_con_la_partida_y_el_cable_se_escribe_una_vez_con_su_metrado(self):
        d = Liquidado(partidas=[partida('*091240')],
                      materiales=[item('CABLE N2XY', 9, codigo='5031165')])
        lineas = [l for l in lineas_del_cuaderno(d) if 'Subida' in l]
        self.assertEqual(lineas, ['Se instaló Subida AP N2XY 2-1x6 (9 metros)'])

    # ── Ferretería ──────────────────────────────────────────────────────────

    def test_mensula_doble(self):
        d = Liquidado(partidas=[partida('*098670')])
        self.assertIn('Se instaló ménsula doble de madera', lineas_del_cuaderno(d))

    def test_diagonal_con_su_cantidad(self):
        d = Liquidado(partidas=[partida('*090060', 2)])
        self.assertIn('Se instaló diagonal de acero (2)', lineas_del_cuaderno(d))

    def test_abrazadera_de_cuatro_pernos_con_su_cantidad(self):
        d = Liquidado(partidas=[partida('*090065', 3)])
        self.assertIn('Se instaló abrazadera de 4 pernos (3)',
                      lineas_del_cuaderno(d))

    # ── Retenidas: van por el tipo de trabajo marcado ───────────────────────

    def test_retenida_tipo_y(self):
        d = Liquidado(tipos=['Retenida Tipo "Y"'])
        self.assertIn('Se instaló retenida tipo "Y"', lineas_del_cuaderno(d))

    def test_retenida_violin(self):
        d = Liquidado(tipos=['Retenida Violin'])
        self.assertIn('Se instaló retenida tipo violín', lineas_del_cuaderno(d))

    def test_retenida_simple(self):
        d = Liquidado(tipos=['Retenida simple'])
        self.assertIn('Se instaló retenida tipo simple', lineas_del_cuaderno(d))

    def test_dos_retenidas_dos_renglones(self):
        d = Liquidado(tipos=['Retenida simple', 'Retenida Violin'])
        lineas = lineas_del_cuaderno(d)
        self.assertIn('Se instaló retenida tipo simple', lineas)
        self.assertIn('Se instaló retenida tipo violín', lineas)

    def test_otro_tipo_de_trabajo_no_escribe_retenida(self):
        d = Liquidado(tipos=['Alumbrado cabria', 'Poste cabria'])
        self.assertFalse(any('retenida' in l for l in lineas_del_cuaderno(d)))

    # ── Cables ──────────────────────────────────────────────────────────────

    def test_los_cables_trasladados_van_uno_por_uno(self):
        d = Liquidado(elementos_plano=[
            cable('T', 25, 'CAAIS 3x16'),
            cable('T', 40, 'CAAIS 3x70'),
        ])
        lineas = lineas_del_cuaderno(d)
        self.assertIn('Se trasladó CAAIS 3x16 25 metros', lineas)
        self.assertIn('Se trasladó CAAIS 3x70 40 metros', lineas)

    def test_lo_instalado_o_existente_no_se_escribe(self):
        d = Liquidado(elementos_plano=[cable('I'), cable('E'), cable('R')])
        self.assertEqual(len(lineas_del_cuaderno(d)), 1)

    def test_cables_de_comunicacion_una_sola_vez(self):
        d = Liquidado(elementos_plano=[cable('C', 0, ''), cable('C', 0, '')])
        lineas = lineas_del_cuaderno(d)
        self.assertEqual(lineas.count('Se trasladó cables de comunicación'), 1)

    # ── Arrastre ────────────────────────────────────────────────────────────

    def test_arrastre_en_pendiente(self):
        d = Liquidado(elementos_plano=[cable('A', 20, '', pendiente=True)])
        self.assertIn('Se realizó arrastre de poste 20 metros en zona de '
                      'pendiente mayor a 30° o escalera', lineas_del_cuaderno(d))

    def test_arrastre_en_plano(self):
        d = Liquidado(elementos_plano=[cable('A', 35, '', pendiente=False)])
        self.assertIn('Se realizó arrastre de poste 35 metros en plano',
                      lineas_del_cuaderno(d))

    def test_los_tramos_de_la_misma_zona_se_suman_en_un_renglon(self):
        """El plano parte el recorrido en pedazos; el cuaderno no."""
        d = Liquidado(elementos_plano=[
            cable('A', 20, '', pendiente=True),
            cable('A', 20, '', pendiente=True),
            cable('A', 30, '', pendiente=True),
            cable('A', 25, '', pendiente=True),
        ])
        arrastres = [l for l in lineas_del_cuaderno(d) if 'arrastre' in l]
        self.assertEqual(arrastres, ['Se realizó arrastre de poste 95 metros '
                                     'en zona de pendiente mayor a 30° o escalera'])

    def test_cada_zona_tiene_su_renglon(self):
        d = Liquidado(elementos_plano=[
            cable('A', 20, '', pendiente=True),
            cable('A', 15, '', pendiente=True),
            cable('A', 40, '', pendiente=False),
        ])
        arrastres = [l for l in lineas_del_cuaderno(d) if 'arrastre' in l]
        self.assertEqual(arrastres, [
            'Se realizó arrastre de poste 35 metros en zona de pendiente '
            'mayor a 30° o escalera',
            'Se realizó arrastre de poste 40 metros en plano',
        ])

    def test_arrastre_sin_responder_se_toma_como_plano(self):
        d = Liquidado(elementos_plano=[cable('A', 12, '')])
        self.assertIn('Se realizó arrastre de poste 12 metros en plano',
                      lineas_del_cuaderno(d))

    # ── Acarreo ─────────────────────────────────────────────────────────────

    def test_el_acarreo_multiplica_el_tramo_por_seis_viajes(self):
        d = Liquidado(partidas=[partida('*090633', 100),
                                partida('*090630', 40),
                                partida('*090632', 25)])
        self.assertIn(
            'Se realizó acarreo de equipos y herramientas por un tramo de 65 '
            'metros por 6 viajes dando en total 390 metros para lo cual se '
            'utilizó la vía más corta', lineas_del_cuaderno(d))

    def test_sin_acarreo_liquidado_no_se_escribe(self):
        d = Liquidado(partidas=[partida('*090630', 40)])
        self.assertFalse(any('acarreo' in l for l in lineas_del_cuaderno(d)))

    # ── Vereda ──────────────────────────────────────────────────────────────

    def test_un_renglon_por_pano_reparado(self):
        d = Liquidado(elementos_plano=[vereda(2, 1.5), vereda(3, 1)])
        lineas = lineas_del_cuaderno(d)
        self.assertIn('Se reparó vereda de 10cm 2 x 1.5 = 3', lineas)
        self.assertIn('Se reparó vereda de 10cm 3 x 1 = 3', lineas)

    def test_un_pano_sin_medidas_no_se_escribe(self):
        d = Liquidado(elementos_plano=[vereda(0, 0)])
        self.assertEqual(len(lineas_del_cuaderno(d)), 1)


    # ── Los postes que salieron ─────────────────────────────────────────────

    def test_poner_un_poste_ya_no_significa_que_salio_el_viejo(self):
        """Antes bastaba con haber consumido un poste del catálogo. En una
        reforma se instala poste sin retirar el que estaba, así que el retiro
        pasó a decirlo solo su partida."""
        d = Liquidado(materiales=[item('POSTE', codigo='5331616')],
                      postes=['900000123'])
        self.assertFalse(any('Se retiró poste' in l for l in lineas_del_cuaderno(d)))

    def test_el_retiro_lo_dice_su_partida(self):
        for codigo in PARTIDAS_RETIRO_POSTE:
            with self.subTest(codigo):
                d = Liquidado(partidas=[partida(codigo)], postes=['900000123'])
                self.assertIn('Se retiró poste 900000123',
                              lineas_del_cuaderno(d))

    def test_sin_partida_de_retiro_no_se_escribe(self):
        d = Liquidado(postes=['900000123'])
        self.assertFalse(any('Se retiró poste' in l for l in lineas_del_cuaderno(d)))

    # ── Lo instalado en la acometida ────────────────────────────────────────

    def test_caja_de_distribucion_por_su_partida(self):
        d = Liquidado(partidas=[partida('*093242')])
        self.assertIn('Se instaló caja de distribución', lineas_del_cuaderno(d))

    def test_corona_con_ganchos_por_su_partida(self):
        d = Liquidado(partidas=[partida('*093045')])
        self.assertIn('Se instaló abrazadera tipo corona con ganchos',
                      lineas_del_cuaderno(d))

    # ── Lo que se bajó, según el recupero ───────────────────────────────────

    def test_la_luminaria_retirada_con_sus_variantes(self):
        d = Liquidado(recuperos=[item('LUMINARIA DE 150 W'),
                                 item('FALORA COMPLETA', 2)])
        self.assertIn('Se retiró luminaria: LUMINARIA DE 150 W (1), '
                      'FALORA COMPLETA (2)', lineas_del_cuaderno(d))

    def test_el_pastoral_retirado(self):
        d = Liquidado(recuperos=[item('PASTORAL JP')])
        self.assertIn('Se retiró pastoral: PASTORAL JP (1)',
                      lineas_del_cuaderno(d))

    def test_la_abrazadera_para_pastoral_no_es_un_pastoral(self):
        d = Liquidado(recuperos=[item('ABRAZADERA PARA PASTORAL')])
        self.assertFalse(any('Se retiró pastoral' in l
                             for l in lineas_del_cuaderno(d)))

    def test_la_caja_y_la_corona_retiradas(self):
        d = Liquidado(recuperos=[item('CAJA DE DISTRIBUCION'),
                                 item('CORONA 4 GANCHOS')])
        lineas = lineas_del_cuaderno(d)
        self.assertIn('Se retiró caja de distribución: CAJA DE DISTRIBUCION (1)',
                      lineas)
        self.assertIn('Se retiró abrazadera tipo corona con ganchos: '
                      'CORONA 4 GANCHOS (1)', lineas)

    # ── Suministros trasladados ─────────────────────────────────────────────

    def test_suministros_trasladados_con_sus_numeros(self):
        d = Liquidado(partidas=[partida('*093081', 2)],
                      conexiones=['1234567', '7654321'])
        self.assertIn('Se realizó traslado de 2 suministros: 1234567; 7654321',
                      lineas_del_cuaderno(d))

    def test_un_solo_suministro_va_en_singular(self):
        d = Liquidado(partidas=[partida('*093081', 1)])
        self.assertIn('Se realizó traslado de 1 suministro',
                      lineas_del_cuaderno(d))

    def test_sin_traslado_de_acometida_no_se_escribe(self):
        self.assertFalse(any('suministro' in l
                             for l in lineas_del_cuaderno(Liquidado())))

    # ── Una SST con varios postes ───────────────────────────────────────────

    def test_con_un_solo_poste_el_cuaderno_no_cambia(self):
        """Lo que ya funcionaba tiene que salir idéntico."""
        d = Liquidado(
            materiales=[item('POSTE', codigo='5331616')],
            postes=['771000100'],
            por_poste=[poste_liquidado('771000100',
                             materiales=[item('POSTE', codigo='5331616')])])
        lineas = lineas_del_cuaderno(d)
        self.assertIn('Se instaló PRFV 9/200', lineas)
        self.assertFalse(any(l.startswith('Poste 771000100:') for l in lineas))

    def test_con_varios_postes_cada_uno_lleva_su_encabezado(self):
        d = Liquidado(
            materiales=[item('POSTE', codigo='5331616', cantidad=2)],
            partidas=[partida('*091320', 2)],
            por_poste=[
                poste_liquidado('771000100',
                      materiales=[item('POSTE', codigo='5331616')],
                      partidas=[partida('*091320')]),
                poste_liquidado('771000101',
                      materiales=[item('POSTE', codigo='5331596')],
                      partidas=[partida('*091320')]),
            ])
        lineas = lineas_del_cuaderno(d)
        self.assertIn('Poste 771000100:', lineas)
        self.assertIn('Poste 771000101:', lineas)
        # Cada uno con el suyo: el de 9 metros y el de 7,5.
        self.assertIn('Se instaló PRFV 9/200', lineas)
        self.assertIn('Se instaló PRFV 7/150', lineas)

    def test_cada_poste_dice_lo_suyo_y_en_su_bloque(self):
        d = Liquidado(
            por_poste=[
                poste_liquidado('771000100', partidas=[partida('*098670')]),
                poste_liquidado('771000101', partidas=[partida('*090060', 2)]),
            ])
        lineas = lineas_del_cuaderno(d)
        primero = lineas.index('Poste 771000100:')
        segundo = lineas.index('Poste 771000101:')
        mensula = lineas.index('Se instaló ménsula doble de madera')
        diagonal = lineas.index('Se instaló diagonal de acero (2)')
        self.assertTrue(primero < mensula < segundo < diagonal)

    def test_lo_del_plano_va_una_sola_vez_al_final(self):
        """El plano es de la SST: si se repitiera por poste, se cobraría dos
        veces y el cuaderno lo diría dos veces."""
        d = Liquidado(
            partidas=[partida('*090633', 100), partida('*090630', 40)],
            elementos_plano=[cable('A', 40, '', pendiente=False),
                             vereda(2, 1.5)],
            por_poste=[
                poste_liquidado('771000100', partidas=[partida('*098670')]),
                poste_liquidado('771000101', partidas=[partida('*090060')]),
            ])
        lineas = lineas_del_cuaderno(d)
        arrastres = [l for l in lineas if 'arrastre' in l]
        veredas = [l for l in lineas if 'vereda' in l]
        self.assertEqual(len(arrastres), 1)
        self.assertEqual(len(veredas), 1)
        # Y después del detalle de los postes.
        self.assertGreater(lineas.index(arrastres[0]),
                           lineas.index('Poste 771000101:'))

    def test_un_poste_sin_nada_no_ensucia_el_cuaderno(self):
        d = Liquidado(por_poste=[
            poste_liquidado('771000100', partidas=[partida('*098670')]),
            poste_liquidado('771000101'),
        ])
        lineas = lineas_del_cuaderno(d)
        self.assertIn('Poste 771000100:', lineas)
        self.assertNotIn('Poste 771000101:', lineas)

    def test_el_retiro_del_poste_va_en_su_bloque(self):
        d = Liquidado(por_poste=[
            poste_liquidado('771000100', partidas=[partida('*090470')]),
            poste_liquidado('771000101', partidas=[partida('*090471')]),
        ])
        lineas = lineas_del_cuaderno(d)
        self.assertIn('Se retiró poste 771000100', lineas)
        self.assertIn('Se retiró poste 771000101', lineas)

    def test_solo_se_nombra_el_poste_que_de_verdad_salio(self):
        """El caso que reportó el usuario: en una reforma el cuaderno decía
        "Se retiró poste Poste 01" en puntos donde no se retiró nada."""
        d = Liquidado(por_poste=[
            poste_liquidado('Poste 01', partidas=[partida('*098670')]),
            poste_liquidado('Poste 02', partidas=[partida('*090470')]),
            poste_liquidado('Poste 03', materiales=[item('POSTE', codigo='5331616')]),
        ])
        retiros = [l for l in lineas_del_cuaderno(d) if 'Se retiró poste' in l]
        self.assertEqual(retiros, ['Se retiró poste Poste 02'])

    # ── El orden de la obra ─────────────────────────────────────────────────

    def test_el_poste_va_antes_que_la_vereda(self):
        d = Liquidado(materiales=[item('POSTE', codigo='5331616')],
                      elementos_plano=[vereda(2, 1)])
        lineas = lineas_del_cuaderno(d)
        self.assertLess(lineas.index('Se instaló PRFV 9/200'),
                        lineas.index('Se reparó vereda de 10cm 2 x 1 = 2'))


class DescargasTests(BaseAPITestCase):

    def setUp(self):
        super().setUp()
        self.actividad = Actividad.objects.create(
            nombre='Cambio de poste inacc. cabria aereo')
        self.sst = SST.objects.create(
            sst='5001', codigo='SST-5001', empresa=self.empresa, distrito='SURCO',
            actividad=self.actividad, fecha_ejecucion=date(2026, 9, 10),
            hora_ejecucion=time(9, 30))
        self.poste = Suministro.objects.create(numero_suministro='P-7788')
        SSTSuministro.objects.create(sst=self.sst, suministro=self.poste)

        self.poste_prfv = Material.objects.create(
            matricula='5331616', descripcion='POSTE POLIESTER 9 / 200', precio='42')
        self.propio = Material.objects.create(
            matricula='9999999', descripcion='MATERIAL QUE NO ESTA', precio='3')
        self.inspeccion = ManoDeObra.objects.create(
            partida='*094395', descripcion='Inspección previa', precio='53.30')
        self.nueva = ManoDeObra.objects.create(
            partida='*777777', descripcion='Partida nueva', precio='5')

    def liquidar(self, sst=None, suministro=None):
        liq = LiquidacionSuministro.objects.create(
            suministro=suministro or self.poste, sst_externo=(sst or self.sst).codigo,
            usuario=self.capataz,
            tipo_trabajo=TipoTrabajo.objects.get_or_create(nombre='Poste cabria')[0])
        LiquidacionPartida.objects.create(
            liquidacion=liq, mano_de_obra=self.inspeccion, cantidad=1)
        LiquidacionPartida.objects.create(
            liquidacion=liq, mano_de_obra=self.nueva, cantidad=3)
        for material, cantidad in ((self.poste_prfv, 1), (self.propio, 4)):
            ConsumoMaterialSuministro.objects.create(
                liquidacion=liq, suministro=suministro or self.poste,
                material=material, usuario=self.capataz, cantidad=cantidad)
        return liq

    def get(self, ruta, codigo=None):
        self.auth(self.capataz)
        return self.client.get(ruta, {'sst': codigo or self.sst.codigo})

    # ── Cuaderno ────────────────────────────────────────────────────────────

    def test_cuaderno_en_pdf(self):
        self.liquidar()
        r = self.get('/api/liquidaciones/cuaderno_obra/')
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b'%PDF'))
        self.assertIn('cuaderno_obra_SST-5001.pdf', r['Content-Disposition'])

    def test_sin_liquidacion_no_hay_cuaderno_ni_numero(self):
        r = self.get('/api/liquidaciones/cuaderno_obra/')
        self.assertEqual(r.status_code, 400)
        self.assertIsInstance(r.data['detail'], str)
        self.assertFalse(CuadernoObra.objects.exists())

    def test_el_numero_se_conserva_y_el_siguiente_es_correlativo(self):
        self.liquidar()
        self.get('/api/liquidaciones/cuaderno_obra/')
        self.get('/api/liquidaciones/cuaderno_obra/')
        otra = SST.objects.create(sst='5002', codigo='SST-5002',
                                  empresa=self.empresa, distrito='LIMA')
        otro_poste = Suministro.objects.create(numero_suministro='P-9900')
        SSTSuministro.objects.create(sst=otra, suministro=otro_poste)
        self.liquidar(sst=otra, suministro=otro_poste)
        self.get('/api/liquidaciones/cuaderno_obra/', 'SST-5002')
        self.assertEqual(
            list(CuadernoObra.objects.order_by('numero')
                 .values_list('sst_codigo', 'numero')),
            [('SST-5001', 1), ('SST-5002', 2)])

    def test_redacta_con_el_plano_y_el_recupero(self):
        from ..cuaderno_obra import reunir
        self.liquidar()
        PlanoSST.objects.create(
            empresa=self.empresa, sst_codigo=self.sst.codigo, usuario=self.capataz,
            elementos=[{'assetId': 'poste_nuevo', 'codigo': '111222333'},
                       {'tipo': 'cable', 'estado': 'A', 'metros': 20}])
        lum = Recupero.objects.create(matricula='REC-036', descripcion='LUMINARIA DE 150 W')
        SuministroRecupero.objects.create(
            suministro=self.poste, recupero=lum, cantidad=1, fecha=date.today())
        lineas = lineas_del_cuaderno(reunir(self.sst))
        self.assertIn('Se instaló PRFV 9/200 con código 111222333', lineas)
        self.assertIn('Se realizó arrastre de poste 20 metros en plano', lineas)

    # ── Excel ───────────────────────────────────────────────────────────────

    def test_excel_con_las_cantidades_en_su_celda(self):
        self.liquidar()
        r = self.get('/api/liquidaciones/excel/')
        self.assertEqual(r.status_code, 200)
        self.assertIn('liquidacion_SST-5001.xlsx', r['Content-Disposition'])
        wb = load_workbook(io.BytesIO(r.content))

        car = wb.worksheets[0]
        self.assertEqual(car['C8'].value, 'SST-5001')
        self.assertEqual(car['C10'].value, 'TECSUR')
        self.assertEqual(car['H18'].value, self.capataz.nombre)

        mat = wb['MATERIAL']
        fila_poste = next(f for f in range(14, 171) if mat[f'A{f}'].value == 5331616)
        self.assertEqual(mat[f'AV{fila_poste}'].value, 1)
        fila_propio = next(f for f in range(14, 171) if mat[f'A{f}'].value == 9999999)
        self.assertGreater(fila_propio, 100)
        self.assertEqual(mat[f'C{fila_propio}'].value, 'MATERIAL QUE NO ESTA')
        self.assertEqual(mat[f'AV{fila_propio}'].value, 4)

        mo = wb['MANO DE OBRA']
        fila_insp = next(f for f in range(7, 113) if mo[f'A{f}'].value == '*094395')
        self.assertEqual(mo[f'F{fila_insp}'].value, 1)
        fila_nueva = next(f for f in range(7, 113) if mo[f'A{f}'].value == '*777777')
        self.assertEqual(mo[f'F{fila_nueva}'].value, 3)
        # La fila nueva se suma sola. En qué columna cae y hasta dónde llega
        # lo manda la plantilla: si se le insertan columnas de postes, la
        # suma se corre y se estira con ellas.
        # Se leen de la plantilla y no del archivo generado: ahí los P1, P2...
        # ya fueron reemplazados por el número de cada poste.
        molde = load_workbook(PLANTILLA)['MANO DE OBRA']
        postes = _columnas_de_postes(molde, MO_ENCABEZADO, MO_POSTES[0])
        suma = _columna_con(molde, MO_ENCABEZADO, 'Cant.')
        self.assertEqual(
            mo[f'{suma}{fila_nueva}'].value,
            f'=SUM({postes[0]}{fila_nueva}:{postes[-1]}{fila_nueva})')

    def test_excel_cables_hasta_35_y_veredas_del_plano(self):
        self.liquidar()
        PlanoSST.objects.create(
            empresa=self.empresa, sst_codigo=self.sst.codigo, usuario=self.capataz,
            elementos=[
                {'tipo': 'cable', 'estado': 'T', 'descripcion': 'Caais 3x35+1x16', 'metros': 26},
                {'tipo': 'cable', 'estado': 'T', 'descripcion': 'Caais 3x70', 'metros': 25},
                {'tipo': 'cable', 'estado': 'I', 'descripcion': 'Caais 3x35', 'metros': 9},
                {'tipo': 'cable', 'estado': 'T', 'descripcion': 'Caais 2x16', 'metros': 15},
                {'tipo': 'cable', 'estado': 'A', 'metros': 30},
                {'tipo': 'vereda', 'largo': 2, 'ancho': 1.5},
                {'tipo': 'vereda', 'largo': 3, 'ancho': 0.8},
            ])
        wb = load_workbook(io.BytesIO(self.get('/api/liquidaciones/excel/').content))
        cab = wb['Cables']
        # Solo lo trasladado de hasta 35 mm2, un tramo por vano.
        self.assertEqual([cab[f'{c}53'].value for c in 'IJKLMN'],
                         [26, 15, None, None, None, None])
        self.assertEqual(cab['O53'].value, '=+SUM(I53:N53)')
        ver = wb['Vereda']
        self.assertEqual((ver['C5'].value, ver['D5'].value), (2, 1.5))
        self.assertEqual((ver['C6'].value, ver['D6'].value), (3, 0.8))
        self.assertIsNone(ver['C7'].value)

    def test_mas_de_seis_tramos_el_ultimo_vano_se_lleva_el_resto(self):
        from ..excel_liquidacion import _llenar_cables
        from openpyxl import Workbook
        ws = Workbook().active
        _llenar_cables(ws, [{'tipo': 'cable', 'estado': 'T', 'descripcion': 'Caais 3x16',
                             'metros': m} for m in (10, 20, 30, 40, 5, 7, 8)])
        self.assertEqual([ws[f'{c}53'].value for c in 'IJKLMN'], [10, 20, 30, 40, 5, 15])

    # ── Una columna por punto de trabajo ────────────────────────────────────

    def sst_de_reforma(self):
        """Una SST de reforma con dos postes, cada uno con lo suyo."""
        reforma = Actividad.objects.create(
            nombre='Reforma - cambio de poste inacc. aereo - cabria',
            varios_postes=True)
        self.sst.actividad = reforma
        self.sst.save(update_fields=['actividad'])
        tipo = TipoTrabajo.objects.get_or_create(nombre='Poste cabria')[0]
        ActividadTipoTrabajo.objects.get_or_create(
            actividad=reforma, tipo_trabajo=tipo)
        otro = Suministro.objects.create(numero_suministro='P-7789')
        SSTSuministro.objects.create(sst=self.sst, suministro=otro)
        return tipo, otro

    def liquidar_poste(self, tipo, poste, partida, cantidad, material=None,
                       cantidad_material=0):
        liq = LiquidacionSuministro.objects.create(
            suministro=poste, sst_externo=self.sst.codigo,
            usuario=self.capataz, tipo_trabajo=tipo)
        LiquidacionPartida.objects.create(
            liquidacion=liq, mano_de_obra=partida, cantidad=cantidad)
        if material is not None:
            ConsumoMaterialSuministro.objects.create(
                liquidacion=liq, suministro=poste, material=material,
                usuario=self.capataz, cantidad=cantidad_material)
        return liq

    def columnas(self, hoja, filas, columnas):
        wb = load_workbook(io.BytesIO(
            self.get('/api/liquidaciones/excel/').content))
        ws = wb[hoja]
        return {c: [ws[f'{c}{f}'].value for f in filas] for c in columnas}

    def test_con_un_solo_poste_todo_va_en_p1(self):
        """Lo que ya funcionaba tiene que salir idéntico."""
        self.liquidar()
        wb = load_workbook(io.BytesIO(
            self.get('/api/liquidaciones/excel/').content))
        ws = wb['MANO DE OBRA']
        fila = [f for f in range(7, 113)
                if str(ws[f'A{f}'].value or '').strip() == '*094395'][0]
        self.assertEqual(ws[f'F{fila}'].value, 1)
        self.assertIsNone(ws[f'G{fila}'].value)

    def test_cada_poste_tiene_su_columna_en_mano_de_obra(self):
        tipo, otro = self.sst_de_reforma()
        self.liquidar_poste(tipo, self.poste, self.inspeccion, 1)
        self.liquidar_poste(tipo, otro, self.inspeccion, 2)
        wb = load_workbook(io.BytesIO(
            self.get('/api/liquidaciones/excel/').content))
        ws = wb['MANO DE OBRA']
        fila = [f for f in range(7, 113)
                if str(ws[f'A{f}'].value or '').strip() == '*094395'][0]
        self.assertEqual(ws[f'F{fila}'].value, 1)
        self.assertEqual(ws[f'G{fila}'].value, 2)

    def test_cada_poste_tiene_su_columna_en_material(self):
        tipo, otro = self.sst_de_reforma()
        self.liquidar_poste(tipo, self.poste, self.inspeccion, 1,
                            material=self.poste_prfv, cantidad_material=1)
        self.liquidar_poste(tipo, otro, self.inspeccion, 1,
                            material=self.poste_prfv, cantidad_material=3)
        wb = load_workbook(io.BytesIO(
            self.get('/api/liquidaciones/excel/').content))
        ws = wb['MATERIAL']
        fila = [f for f in range(14, 171)
                if str(ws[f'A{f}'].value or '').strip() == '5331616'][0]
        self.assertEqual(ws[f'AV{fila}'].value, 1)
        self.assertEqual(ws[f'AW{fila}'].value, 3)

    def test_la_suma_de_las_columnas_es_el_total(self):
        tipo, otro = self.sst_de_reforma()
        self.liquidar_poste(tipo, self.poste, self.inspeccion, 1)
        self.liquidar_poste(tipo, otro, self.inspeccion, 2)
        wb = load_workbook(io.BytesIO(
            self.get('/api/liquidaciones/excel/').content))
        ws = wb['MANO DE OBRA']
        fila = [f for f in range(7, 113)
                if str(ws[f'A{f}'].value or '').strip() == '*094395'][0]
        # La plantilla suma F:H en I; aquí se comprueba lo que se escribió.
        escrito = [ws[f'{c}{fila}'].value or 0 for c in 'FGH']
        self.assertEqual(sum(escrito), 3)

    def test_un_poste_que_no_liquido_esa_partida_deja_su_columna_vacia(self):
        tipo, otro = self.sst_de_reforma()
        self.liquidar_poste(tipo, self.poste, self.inspeccion, 1)
        self.liquidar_poste(tipo, otro, self.nueva, 5)
        wb = load_workbook(io.BytesIO(
            self.get('/api/liquidaciones/excel/').content))
        ws = wb['MANO DE OBRA']
        fila = [f for f in range(7, 113)
                if str(ws[f'A{f}'].value or '').strip() == '*094395'][0]
        self.assertEqual(ws[f'F{fila}'].value, 1)
        self.assertIsNone(ws[f'G{fila}'].value)

    def test_las_columnas_llevan_el_numero_de_cada_poste(self):
        """Los P1, P2, P3 de la plantilla dicen qué poste es cada columna."""
        from ..excel_liquidacion import MAT_ENCABEZADO, MO_ENCABEZADO
        tipo, otro = self.sst_de_reforma()
        self.liquidar_poste(tipo, self.poste, self.inspeccion, 1)
        self.liquidar_poste(tipo, otro, self.inspeccion, 2)
        wb = load_workbook(io.BytesIO(
            self.get('/api/liquidaciones/excel/').content))
        mat, mo = wb['MATERIAL'], wb['MANO DE OBRA']
        self.assertEqual(mat[f'AV{MAT_ENCABEZADO}'].value,
                         self.poste.numero_suministro)
        self.assertEqual(mat[f'AW{MAT_ENCABEZADO}'].value, 'P-7789')
        self.assertEqual(mo[f'F{MO_ENCABEZADO}'].value,
                         self.poste.numero_suministro)
        self.assertEqual(mo[f'G{MO_ENCABEZADO}'].value, 'P-7789')

    def test_el_orden_de_las_columnas_es_el_de_grabado(self):
        """El usuario las quiere en el orden en que se graban los postes."""
        from ..excel_liquidacion import MO_ENCABEZADO
        tipo, otro = self.sst_de_reforma()
        self.liquidar_poste(tipo, otro, self.inspeccion, 2)
        self.liquidar_poste(tipo, self.poste, self.inspeccion, 1)
        wb = load_workbook(io.BytesIO(
            self.get('/api/liquidaciones/excel/').content))
        ws = wb['MANO DE OBRA']
        self.assertEqual(ws[f'F{MO_ENCABEZADO}'].value, 'P-7789')
        self.assertEqual(ws[f'G{MO_ENCABEZADO}'].value,
                         self.poste.numero_suministro)

    # ── Hoja Traslado - Acarreo ─────────────────────────────────────────────

    def plano_con_arrastre(self, *metros):
        self.liquidar()
        PlanoSST.objects.create(
            empresa=self.empresa, sst_codigo=self.sst.codigo,
            usuario=self.capataz,
            elementos=[{'tipo': 'cable', 'estado': 'A', 'metros': m}
                       for m in metros])
        wb = load_workbook(io.BytesIO(
            self.get('/api/liquidaciones/excel/').content))
        return wb['Traslado - Acarreo']

    def fila(self, ws, numero):
        return [ws[f'{c}{numero}'].value for c in 'CDEFGHIJK']

    def test_cada_tramo_de_arrastre_va_en_su_columna(self):
        ws = self.plano_con_arrastre(40, 25, 55)
        self.assertEqual(self.fila(ws, 4),
                         [40, 25, 55, None, None, None, None, None, None])

    def test_el_bloque_de_acarreo_se_llena_solo_desde_la_fila_4(self):
        """La plantilla ya trae C19 = C4, C20 = D4, ... : no se tocan."""
        ws = self.plano_con_arrastre(40, 25, 55)
        self.assertEqual([ws[f'C{f}'].value for f in (19, 20, 21)],
                         ['=+C4', '=+D4', '=+E4'])
        self.assertEqual(ws['M31'].value, 6)

    def test_hasta_cien_metros_no_se_cobra_traslado(self):
        ws = self.plano_con_arrastre(40, 25)
        self.assertEqual(self.fila(ws, 5), [None] * 9)

    def test_pasando_los_cien_se_repiten_los_tramos_con_su_descuento(self):
        ws = self.plano_con_arrastre(40, 25, 55)
        self.assertEqual(self.fila(ws, 5),
                         [40, 25, 55, None, None, None, None, None, -100])

    def test_justo_cien_metros_todavia_esta_incluido(self):
        ws = self.plano_con_arrastre(60, 40)
        self.assertEqual(self.fila(ws, 5), [None] * 9)

    def test_con_mas_tramos_que_columnas_la_ultima_se_lleva_el_resto(self):
        ws = self.plano_con_arrastre(*([10] * 11))
        fila4 = self.fila(ws, 4)
        self.assertEqual(fila4[:8], [10] * 8)
        self.assertEqual(fila4[8], 30)          # los tres que sobran
        self.assertEqual(sum(v for v in fila4 if v), 110)

    def test_en_la_fila_de_cobro_el_descuento_no_pisa_un_tramo(self):
        """K5 es del descuento, así que esa fila tiene una columna menos."""
        ws = self.plano_con_arrastre(*([20] * 9))
        fila5 = self.fila(ws, 5)
        self.assertEqual(fila5[:7], [20] * 7)
        self.assertEqual(fila5[7], 40)          # los dos que sobran
        self.assertEqual(fila5[8], -100)

    def test_sin_arrastre_la_hoja_queda_como_estaba(self):
        ws = self.plano_con_arrastre()
        self.assertEqual(self.fila(ws, 4), [None] * 9)
        self.assertEqual(self.fila(ws, 5), [None] * 9)

    def test_los_cables_trasladados_no_son_arrastre(self):
        self.liquidar()
        PlanoSST.objects.create(
            empresa=self.empresa, sst_codigo=self.sst.codigo,
            usuario=self.capataz,
            elementos=[{'tipo': 'cable', 'estado': 'T',
                        'descripcion': 'Caais 3x16', 'metros': 120}])
        wb = load_workbook(io.BytesIO(
            self.get('/api/liquidaciones/excel/').content))
        self.assertEqual(self.fila(wb['Traslado - Acarreo'], 4), [None] * 9)

    def test_el_cable_instalado_va_en_la_fila_de_instalacion(self):
        """La plantilla rotula la fila 52 como instalación y la 53 traslado."""
        self.liquidar()
        PlanoSST.objects.create(
            empresa=self.empresa, sst_codigo=self.sst.codigo,
            usuario=self.capataz,
            elementos=[
                {'tipo': 'cable', 'estado': 'I', 'descripcion': 'Caais 3x16',
                 'metros': 40},
                {'tipo': 'cable', 'estado': 'I', 'descripcion': 'Caais 2x16',
                 'metros': 15},
                {'tipo': 'cable', 'estado': 'T',
                 'descripcion': 'Caais 3x35+1x16', 'metros': 26},
                {'tipo': 'cable', 'estado': 'I', 'descripcion': 'Caais 3x70',
                 'metros': 90},
            ])
        wb = load_workbook(io.BytesIO(
            self.get('/api/liquidaciones/excel/').content))
        cab = wb['Cables']
        self.assertEqual([cab[f'{c}52'].value for c in 'IJKLMN'],
                         [40, 15, 0, 0, 0, 0])
        # El grueso no es de hasta 35: no entra en esta fila.
        self.assertEqual([cab[f'{c}53'].value for c in 'IJ'], [26, None])

    # ── El Excel cobra lo mismo que el consolidado ──────────────────────────

    def _tipo_poste(self):
        tipo = TipoTrabajo.objects.get_or_create(nombre='Poste cabria')[0]
        ActividadTipoTrabajo.objects.get_or_create(
            actividad=self.actividad, tipo_trabajo=tipo)
        return tipo

    def liquidar_cambio_de_poste(self, acarreo=None, vereda=None):
        """Un cambio de poste con lo que el paquete ya incluye."""
        cambio = ManoDeObra.objects.create(
            partida='*090470', descripcion='CAMBIO DE POSTE CON VEREDA',
            precio='100')
        liq = LiquidacionSuministro.objects.create(
            suministro=self.poste, sst_externo=self.sst.codigo,
            usuario=self.capataz, tipo_trabajo=self._tipo_poste())
        LiquidacionPartida.objects.create(
            liquidacion=liq, mano_de_obra=cambio, cantidad=1)
        if acarreo is not None:
            mo = ManoDeObra.objects.create(
                partida='*090633', descripcion='ACARREO', precio='2')
            LiquidacionPartida.objects.create(
                liquidacion=liq, mano_de_obra=mo, cantidad=acarreo)
        if vereda is not None:
            mo = ManoDeObra.objects.create(
                partida='*091840', descripcion='ROTURA DE VEREDA', precio='24.55')
            LiquidacionPartida.objects.create(
                liquidacion=liq, mano_de_obra=mo, cantidad=vereda)
        return liq

    def mano_de_obra_del_excel(self):
        r = self.get('/api/liquidaciones/excel/')
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content))['MANO DE OBRA']
        filas = {}
        for fila in range(7, 113):
            codigo = ws[f'A{fila}'].value
            if codigo:
                filas[str(codigo).strip()] = ws[f'F{fila}'].value
        return filas

    def test_el_excel_descuenta_lo_que_el_paquete_incluye(self):
        """120 de acarreo con un cambio de poste: 100 van incluidos."""
        self.liquidar_cambio_de_poste(acarreo=120)
        self.assertEqual(self.mano_de_obra_del_excel()['*090633'], 20)

    def test_la_vereda_incluida_sale_en_cero(self):
        """Dos roturas por cambio CON vereda: las dos están incluidas."""
        self.liquidar_cambio_de_poste(vereda=2)
        self.assertEqual(self.mano_de_obra_del_excel()['*091840'], 0)

    def test_lo_que_pasa_de_lo_incluido_si_se_cobra(self):
        self.liquidar_cambio_de_poste(vereda=5)
        self.assertEqual(self.mano_de_obra_del_excel()['*091840'], 3)

    def test_sin_cambio_de_poste_no_se_descuenta_nada(self):
        liq = LiquidacionSuministro.objects.create(
            suministro=self.poste, sst_externo=self.sst.codigo,
            usuario=self.capataz, tipo_trabajo=self._tipo_poste())
        mo = ManoDeObra.objects.create(
            partida='*090633', descripcion='ACARREO', precio='2')
        LiquidacionPartida.objects.create(
            liquidacion=liq, mano_de_obra=mo, cantidad=120)
        self.assertEqual(self.mano_de_obra_del_excel()['*090633'], 120)

    def test_el_excel_dice_lo_mismo_que_el_consolidado(self):
        self.liquidar_cambio_de_poste(acarreo=120, vereda=5)
        excel = self.mano_de_obra_del_excel()
        self.auth(self.capataz)
        r = self.client.get('/api/liquidaciones/consolidado/')
        fila = [x for x in r.data if x['sst'] == self.sst.codigo][0]
        for p in fila['partidas']:
            if p['partida'] in excel:
                self.assertEqual(
                    excel[p['partida']], float(p['cantidad_cobrada']),
                    f'{p["partida"]} no coincide')

    def test_excel_sin_liquidacion(self):
        r = self.get('/api/liquidaciones/excel/')
        self.assertEqual(r.status_code, 400)

    # ── Hora de ejecución ───────────────────────────────────────────────────

    def test_la_hora_se_guarda_con_la_fecha(self):
        self.auth(self.capataz)
        r = self.client.post('/api/ssts/set_fecha_ejecucion/', {
            'sst_codigo': self.sst.codigo, 'fecha': '2026-09-12', 'hora': '14:45',
        }, format='json')
        self.assertEqual(r.status_code, 200, r.data)
        self.sst.refresh_from_db()
        self.assertEqual(self.sst.fecha_ejecucion, date(2026, 9, 12))
        self.assertEqual(self.sst.hora_ejecucion, time(14, 45))

    def test_sin_hora_la_hora_no_se_borra(self):
        self.auth(self.capataz)
        self.client.post('/api/ssts/set_fecha_ejecucion/', {
            'sst_codigo': self.sst.codigo, 'fecha': '2026-09-12'}, format='json')
        self.sst.refresh_from_db()
        self.assertEqual(self.sst.hora_ejecucion, time(9, 30))

    def test_el_comentario_del_tipo_se_guarda_aparte(self):
        self.auth(self.capataz)
        tipo = TipoTrabajo.objects.create(nombre='Conexiones cabria')
        r = self.client.post('/api/liquidaciones/', {
            'suministro': self.poste.pk, 'sst_externo': self.sst.codigo,
            'usuario': self.capataz.pk, 'tipo_trabajo': tipo.pk,
            'observacion': 'general — 1234567', 'comentario': '1234567',
            'partidas': [{'mano_de_obra': self.inspeccion.pk, 'cantidad': 1}],
        }, format='json')
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(LiquidacionSuministro.objects.get(pk=r.data['id_liquidacion'])
                         .comentario, '1234567')


class ColumnasDePostesTests(SimpleTestCase):
    """Cuántas columnas de postes hay lo dice la plantilla, no el código.

    Así, cuando se le insertan columnas en Excel y se rotulan P4, P5..., el
    Excel las usa sin que haya que tocar nada aquí."""

    def hoja(self, encabezados):
        """Una hoja suelta con un encabezado como el de la plantilla."""
        wb = Workbook()
        ws = wb.active
        for columna, valor in encabezados.items():
            ws[f'{columna}1'] = valor
        return ws

    def test_las_encuentra_desde_la_primera(self):
        ws = self.hoja({'E': 'S/,', 'F': 'P1', 'G': 'P2', 'H': 'P3',
                        'I': 'Cant.'})
        self.assertEqual(_columnas_de_postes(ws, 1, 'F'), ('F', 'G', 'H'))

    def test_si_la_plantilla_trae_mas_las_usa_todas(self):
        ws = self.hoja({'F': 'P1', 'G': 'P2', 'H': 'P3', 'I': 'P4',
                        'J': 'P5', 'K': 'P6', 'L': 'Cant.'})
        self.assertEqual(_columnas_de_postes(ws, 1, 'F'),
                         ('F', 'G', 'H', 'I', 'J', 'K'))

    def test_se_corta_donde_deja_de_decir_P(self):
        ws = self.hoja({'F': 'P1', 'G': 'P2', 'H': 'Cant.', 'I': 'P3'})
        self.assertEqual(_columnas_de_postes(ws, 1, 'F'), ('F', 'G'))

    def test_sin_rotulos_queda_la_primera(self):
        self.assertEqual(_columnas_de_postes(self.hoja({}), 1, 'F'), ('F',))

    def test_las_columnas_de_la_derecha_se_buscan_por_su_rotulo(self):
        """Insertar columnas de postes corre Cant., Cant. Final y Total Final."""
        ws = self.hoja({'F': 'P1', 'G': 'P2', 'H': 'P3', 'I': 'P4',
                        'J': 'Cant.', 'K': 'AP/SUBT.', 'M': 'Cant. Final',
                        'N': 'Total Final'})
        self.assertEqual(_columna_con(ws, 1, 'Cant.'), 'J')
        self.assertEqual(_columna_con(ws, 1, 'Cant. Final'), 'M')
        self.assertEqual(_columna_con(ws, 1, 'Total Final'), 'N')
        self.assertIsNone(_columna_con(ws, 1, 'No existe'))

    def test_con_mas_postes_que_columnas_la_ultima_los_nombra_a_todos(self):
        """La última columna se lleva las cantidades de los que sobran, así que
        tiene que decir de quiénes es."""
        ws = self.hoja({'F': 'P1', 'G': 'P2', 'H': 'P3'})
        _rotular_postes(ws, 1, ('F', 'G', 'H'),
                        ['771000100', '771000101', '771000102', '771000103'])
        self.assertEqual(ws['F1'].value, '771000100')
        self.assertEqual(ws['G1'].value, '771000101')
        self.assertEqual(ws['H1'].value, '771000102 + 771000103')

    def test_sin_postes_la_plantilla_queda_como_viene(self):
        ws = self.hoja({'F': 'P1', 'G': 'P2'})
        _rotular_postes(ws, 1, ('F', 'G'), [])
        self.assertEqual(ws['F1'].value, 'P1')
