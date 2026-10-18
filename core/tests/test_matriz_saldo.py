"""La matriz del saldo del camión: una pestaña por inventario, cada una con su
saldo inicial (lo contado en el anterior), cada movimiento por material en
orden de fecha y al final el saldo calculado contra lo contado. La última va
hasta hoy y cierra con el saldo del sistema."""
import io
from datetime import date, timedelta
from decimal import Decimal

from openpyxl import load_workbook

from ..matriz_saldo import ACTUAL, CORTE, ENTRA, INVENTARIO, SALE, matriz
from ..models import (
    AsignacionAgregado, ConsumoMaterialSuministro, DetalleDevolucion,
    DetalleInventario, DetallePedido, Devolucion, Inventario,
    LiquidacionSuministro, Pedido, StockCamion, TipoTrabajo, Usuario,
)
from .base import BaseAPITestCase

URL = '/api/stock-camion/matriz/'


class MatrizSaldoTests(BaseAPITestCase):

    def setUp(self):
        super().setUp()
        self.hoy = date.today()
        # Entran 10 del A por pedido, se liquidan 3 y se devuelven 2: quedan 5.
        self.pedido = Pedido.objects.create(
            camion=self.camion, usuario=self.capataz, estado='aprobado',
            fecha_aprobacion=self.hoy - timedelta(days=5))
        DetallePedido.objects.create(pedido=self.pedido, material=self.material_a,
                                     cantidad_solicitada=10, cantidad_aprobada=10)
        liq = LiquidacionSuministro.objects.create(
            usuario=self.capataz, tipo_trabajo=TipoTrabajo.objects.create(nombre='Cambio'),
            suministro_externo='1014213', sst_externo='5551234')
        ConsumoMaterialSuministro.objects.create(
            liquidacion=liq, suministro_externo='1014213', material=self.material_a,
            usuario=self.capataz, cantidad=3)
        devolucion = Devolucion.objects.create(
            camion=self.camion, usuario=self.capataz, estado='aprobado',
            fecha_aprobacion=self.hoy)
        DetalleDevolucion.objects.create(devolucion=devolucion, material=self.material_a,
                                         cantidad_solicitada=2, cantidad_aprobada=2)
        self.stock = StockCamion.objects.create(camion=self.camion,
                                                material=self.material_a, cantidad=5)

    def inventario(self, dias_atras, contado):
        inv = Inventario.objects.create(camion=self.camion, usuario=self.encargado,
                                        mes=self.hoy.month, anio=self.hoy.year,
                                        estado='cerrado')
        Inventario.objects.filter(pk=inv.pk).update(
            fecha=self.hoy - timedelta(days=dias_atras))
        for material, (teorico, fisico) in contado.items():
            DetalleInventario.objects.create(
                inventario=inv, material=material, cantidad_teorica=teorico,
                cantidad_fisica=fisico, diferencia=fisico - teorico)
        return Inventario.objects.get(pk=inv.pk)

    def test_sin_inventario_hay_una_sola_pestana_hasta_hoy(self):
        [(p, [fila])] = matriz(self.camion)
        self.assertEqual(p.cierre, ACTUAL)
        self.assertEqual(p.nombre, 'Actual')
        self.assertEqual([m.tipo for m in p.movs], ['Pedido', 'Liquidación', 'Devolución'])
        self.assertEqual(p.movs[1].documento, 'SST 5551234 · poste 1014213')
        self.assertEqual([m.clase for m in p.movs], [ENTRA, SALE, SALE])
        self.assertEqual(fila['inicial'], 0)
        self.assertEqual(fila['celdas'], [Decimal(10), Decimal(-3), Decimal(-2)])
        self.assertEqual((fila['entra'], fila['sale'], fila['calculado']), (10, 5, 5))
        self.assertEqual((fila['cierre'], fila['diferencia']), (5, 0))

    def test_cada_inventario_cierra_una_pestana_y_abre_la_siguiente(self):
        # Al inventario de hace 3 días había 10 y se contaron 9: falta 1.
        inv = self.inventario(3, {self.material_a: (10, 9)})
        self.stock.cantidad = 4
        self.stock.save()
        (antes, [fa]), (actual, [fb]) = matriz(self.camion)
        self.assertEqual(antes.cierre, INVENTARIO)
        self.assertEqual(antes.nombre, f'Hasta Inv {inv.pk} ({inv.fecha:%d-%m-%y})')
        self.assertEqual([m.tipo for m in antes.movs], ['Pedido'])
        self.assertEqual((fa['calculado'], fa['cierre'], fa['diferencia']), (10, 9, -1))
        # La siguiente arranca con lo contado.
        self.assertEqual(actual.nombre, f'Actual (desde Inv {inv.pk})')
        self.assertEqual([m.tipo for m in actual.movs], ['Liquidación', 'Devolución'])
        self.assertEqual((fb['inicial'], fb['calculado'], fb['diferencia']), (9, 4, 0))

    def test_lo_que_no_se_conto_sigue_con_su_saldo_calculado(self):
        DetallePedido.objects.create(pedido=self.pedido, material=self.material_b,
                                     cantidad_solicitada=6, cantidad_aprobada=6)
        self.inventario(3, {self.material_a: (10, 10)})
        (antes, filas_antes), (actual, filas) = matriz(self.camion)
        b_antes = next(f for f in filas_antes if f['matricula'] == 'MAT-B')
        self.assertEqual((b_antes['cierre'], b_antes['diferencia']), (6, 0))
        b = next(f for f in filas if f['matricula'] == 'MAT-B')
        self.assertEqual(b['inicial'], 6)

    def test_sin_inventario_por_mas_de_dos_meses_se_parte_por_mes(self):
        self.pedido.fecha_aprobacion = self.hoy - timedelta(days=100)
        self.pedido.save()
        pestanas = matriz(self.camion)
        self.assertGreaterEqual(len(pestanas), 4)
        self.assertTrue(all(p.cierre == CORTE for p, _ in pestanas[:-1]))
        self.assertEqual(pestanas[-1][0].cierre, ACTUAL)
        self.assertTrue(pestanas[0][0].nombre.endswith('sin inventario'))
        # El saldo pasa de una pestaña a la otra.
        [segunda] = pestanas[1][1]
        self.assertEqual(segunda['inicial'], 10)
        [ultima] = pestanas[-1][1]
        self.assertEqual((ultima['calculado'], ultima['diferencia']), (5, 0))

    def test_las_asignaciones_del_mismo_dia_van_en_una_columna(self):
        for material, cantidad in ((self.material_a, 4), (self.material_b, 6),
                                   (self.material_b, 1)):
            AsignacionAgregado.objects.create(
                almacen=self.almacen, camion=self.camion, material=material,
                cantidad=cantidad, usuario=self.capataz, registrado_por=self.encargado,
                origen=AsignacionAgregado.ORIGEN_ALMACEN)
        [(p, filas)] = matriz(self.camion)
        [a] = [m for m in p.movs if m.tipo == 'Asignación']
        self.assertEqual(a.documento.count(','), 2)
        self.assertEqual(a.responsable, 'Capataz Uno')
        self.assertEqual(a.cantidades, {self.material_a.pk: 4, self.material_b.pk: 7})
        self.assertEqual({f['matricula']: f['calculado'] for f in filas},
                         {'MAT-A': 9, 'MAT-B': 7})

    def test_no_cuenta_lo_que_otro_liquido_con_otro_camion(self):
        otro = Usuario.objects.create(nombre='Otro', rol=self.rol_capataz,
                                      empresa=self.empresa, email='o@t.pe', clave='x')
        liq = LiquidacionSuministro.objects.create(
            usuario=otro, tipo_trabajo=TipoTrabajo.objects.create(nombre='Otro'))
        ConsumoMaterialSuministro.objects.create(
            liquidacion=liq, material=self.material_a, usuario=otro, cantidad=7)
        [(_, [fila])] = matriz(self.camion)
        self.assertEqual(fila['calculado'], 5)

    def test_el_almacen_baja_el_excel_con_una_pestana_por_inventario(self):
        inv = self.inventario(3, {self.material_a: (10, 9)})
        self.stock.cantidad = 4
        self.stock.save()
        self.auth(self.encargado)
        r = self.client.get(URL, {'camion': self.camion.pk})
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, [
            f'Hasta Inv {inv.pk} ({inv.fecha:%d-%m-%y})',
            f'Actual (desde Inv {inv.pk})', 'Movimientos'])
        self.assertEqual(wb.active.title, f'Actual (desde Inv {inv.pk})')
        ws = wb.worksheets[0]
        self.assertEqual(ws['A7'].value, 'MAT-A')
        self.assertEqual(ws['C4'].value, 'Saldo inicial')
        self.assertEqual(ws.cell(5, 4).value, 'Pedido')
        self.assertEqual([ws.cell(4, c).value for c in (5, 6, 7, 8, 9)],
                         ['Total entró', 'Total salió', 'Saldo calculado',
                          'Contado en inventario', 'Diferencia (faltante −)'])
        self.assertEqual([ws.cell(7, c).value for c in (3, 4, 7, 8, 9)], [0, 10, 10, 9, -1])
        ws = wb.worksheets[1]
        self.assertEqual([ws.cell(5, c).value for c in (4, 5)], ['Liquidación', 'Devolución'])
        self.assertEqual([ws.cell(7, c).value for c in (3, 4, 5, 8, 9, 10)],
                         [9, -3, -2, 4, 4, 0])
        mov = wb['Movimientos']
        self.assertEqual([c.value for c in mov['B'][1:]],
                         ['Pedido', 'Inventario', 'Liquidación', 'Devolución'])
        self.assertEqual([c.value for c in mov['H'][1:]], [10, 9, 6, 4])

    def test_el_responsable_baja_el_de_su_camion_y_nadie_mas(self):
        self.auth(self.capataz)
        self.assertEqual(self.client.get(URL, {'camion': self.camion.pk}).status_code, 200)
        otro = Usuario.objects.create(nombre='Otro', rol=self.rol_capataz,
                                      empresa=self.empresa, email='o@t.pe', clave='x')
        self.auth(otro)
        self.assertEqual(self.client.get(URL, {'camion': self.camion.pk}).status_code, 403)
        self.assertEqual(self.client.get(URL).status_code, 400)
