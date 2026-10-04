"""
Hoja Vereda del Excel: los paños que se miden al liquidar el reforzamiento
van cada uno al bloque de su espesor.
"""
import io
from datetime import date

from django.test import SimpleTestCase
from openpyxl import load_workbook

from ..excel_liquidacion import generar_excel_liquidacion
from ..models import (SST, Actividad, LiquidacionPartida, LiquidacionSuministro,
                      ManoDeObra, SSTSuministro, Suministro, TipoTrabajo)
from .base import BaseAPITestCase


def hoja(panos, plano=()):
    contenido = generar_excel_liquidacion(
        {'sst': 'SST-1'}, [], [], list(plano), panos_medidos=panos)
    return load_workbook(io.BytesIO(contenido))['Vereda']


def vereda(largo, ancho):
    return {'tipo': 'vereda', 'largo': largo, 'ancho': ancho}


class BloquesTests(SimpleTestCase):

    def test_cada_espesor_a_su_bloque(self):
        ws = hoja({
            'vereda_10': [[1.5, 1.0], [0.8, 0.5]],
            'vereda_15': [[1.2, 0.6]],
            'vereda_20': [[2.0, 1.0]],
            'pista': [[3.0, 1.0]],
            'asfalto': [[1.0, 0.6]],
            'piso_especial': [[1.0, 1.0]],
            'grass': [[2.0, 2.0]],
        })
        self.assertEqual((ws['C5'].value, ws['D5'].value), (1.5, 1.0))
        self.assertEqual((ws['C6'].value, ws['D6'].value), (0.8, 0.5))
        self.assertEqual((ws['C66'].value, ws['D66'].value), (1.2, 0.6))
        # La de 20 cm va en el bloque Sardinel 20cm.
        self.assertEqual((ws['J5'].value, ws['K5'].value), (2.0, 1.0))
        self.assertEqual((ws['C78'].value, ws['D78'].value), (3.0, 1.0))
        self.assertEqual((ws['J78'].value, ws['K78'].value), (1.0, 0.6))
        self.assertEqual((ws['C98'].value, ws['D98'].value), (1.0, 1.0))
        self.assertEqual((ws['C111'].value, ws['D111'].value), (2.0, 2.0))

    def test_el_primer_piso_especial_suma(self):
        ws = hoja({'piso_especial': [[1.0, 1.0]]})
        self.assertEqual(ws['F98'].value, '=+C98*D98')

    def test_lo_que_no_entra_no_pisa_el_bloque_de_abajo(self):
        ws = hoja({'vereda_15': [[1, 1]] * 7})
        self.assertEqual(ws['C70'].value, 1)
        self.assertIsNone(ws['C71'].value)
        self.assertEqual(ws['D71'].value, 'Descuentos')

    def test_los_medidos_siguen_debajo_de_los_del_plano(self):
        ws = hoja({'vereda_10': [[1.5, 1.0]]},
                  plano=[vereda(2, 1), vereda(3, 1)])
        self.assertEqual(ws['C5'].value, 2)
        self.assertEqual(ws['C6'].value, 3)
        self.assertEqual(ws['C7'].value, 1.5)

    def test_sin_medidas_queda_como_antes(self):
        ws = hoja(None, plano=[vereda(2, 1)])
        self.assertEqual(ws['C5'].value, 2)
        self.assertIsNone(ws['J5'].value)
        self.assertIsNone(ws['C66'].value)


class EndpointTests(BaseAPITestCase):

    def test_el_excel_toma_los_panos_de_la_liquidacion(self):
        actividad = Actividad.objects.create(
            nombre='Reforzamiento de poste con vereda', de_encargado=True)
        sst = SST.objects.create(
            sst='9001', codigo='SST-9001', empresa=self.empresa,
            actividad=actividad, fecha_ejecucion=date(2026, 10, 2))
        poste = Suministro.objects.create(numero_suministro='P-9001')
        SSTSuministro.objects.create(sst=sst, suministro=poste)
        liq = LiquidacionSuministro.objects.create(
            suministro=poste, sst_externo='SST-9001', usuario=self.capataz,
            tipo_trabajo=TipoTrabajo.objects.create(
                nombre='Reforzamiento con vereda'),
            medidas={'panos': {'vereda_10': [[1.5, 1.0]],
                               'vereda_20': [[0.8, 0.5]]}})
        LiquidacionPartida.objects.create(
            liquidacion=liq, cantidad=1, mano_de_obra=ManoDeObra.objects.create(
                partida='*094395', descripcion='INSPECCION', precio='53.30'))
        self.auth(self.capataz)
        r = self.client.get('/api/liquidaciones/excel/', {'sst': 'SST-9001'})
        self.assertEqual(r.status_code, 200, r.content[:300])
        ws = load_workbook(io.BytesIO(r.content))['Vereda']
        self.assertEqual((ws['C5'].value, ws['D5'].value), (1.5, 1.0))
        self.assertEqual((ws['J5'].value, ws['K5'].value), (0.8, 0.5))
