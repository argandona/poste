"""Historial de EPP entregado: una fila por trabajador y EPP, una columna por
entrega, con cuántos días le duró cada una.

Pedido del usuario el 2026-10-05: lo ven el SuperAdmin, el Coordinador y el
encargado de almacén, en
la app y en Excel. Un EPP se agrupa por su descripción, no por la talla: si a
alguien le cambian las botas 42 por unas 43, siguen siendo sus botas. Lo que
duró una entrega son los días hasta la siguiente entrega de ese mismo EPP; la
última sigue en uso y cuenta hasta hoy.
"""
import io
from datetime import date

from django.utils import timezone


def armar_historial(empresa_id=None, hoy=None):
    """Las filas de la matriz, ordenadas por trabajador y EPP."""
    from .models import DetallePedidoEPP

    hoy = hoy or timezone.localdate()
    detalles = (DetallePedidoEPP.objects
                .filter(pedido__estado='aprobado', cantidad_aprobada__gt=0)
                .select_related('epp', 'pedido__usuario')
                .order_by('pedido__fecha_aprobacion', 'id_detalle'))
    if empresa_id:
        detalles = detalles.filter(pedido__usuario__empresa_id=empresa_id)

    filas = {}
    for d in detalles:
        u = d.pedido.usuario
        fila = filas.setdefault((u.nombre, u.pk, d.epp.descripcion), {
            'usuario': u.pk,
            'trabajador': u.nombre,
            'dni': u.dni,
            'epp': d.epp.descripcion,
            'unidad': d.epp.unidad,
            'entregas': [],
        })
        fila['entregas'].append({
            'fecha': timezone.localtime(d.pedido.fecha_aprobacion).date(),
            'talla': d.epp.talla,
            'cantidad': d.cantidad_aprobada,
            'tipo': d.pedido.tipo,
            'motivo': d.pedido.motivo,
        })

    for fila in filas.values():
        entregas = fila['entregas']
        for i, e in enumerate(entregas):
            hasta = entregas[i + 1]['fecha'] if i + 1 < len(entregas) else hoy
            e['dias'] = (hasta - e['fecha']).days
            e['en_uso'] = i + 1 == len(entregas)
    return [filas[k] for k in sorted(filas)]


def texto_de_entrega(e):
    """05/10/2026 · T42 · 45 días (en uso)"""
    partes = [e['fecha'].strftime('%d/%m/%Y')]
    if e['talla']:
        partes.append(f"T{e['talla']}")
    if e['cantidad'] != 1:
        partes.append(f"x{e['cantidad'].normalize():f}")
    dias = f"{e['dias']} día{'s' if e['dias'] != 1 else ''}"
    partes.append(f"{dias} (en uso)" if e['en_uso'] else dias)
    return ' · '.join(partes)


def excel_historial(filas, hoy=None):
    """La misma matriz en un .xlsx."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    hoy = hoy or timezone.localdate()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Historial EPP'
    columnas = max((len(f['entregas']) for f in filas), default=1)
    ws.append([f'Historial de EPP entregado al {hoy.strftime("%d/%m/%Y")}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    encabezado = ['Trabajador', 'DNI', 'EPP', 'Entregas'] + [
        f'Entrega {i}' for i in range(1, columnas + 1)]
    ws.append(encabezado)
    gris = PatternFill('solid', fgColor='DDDDDD')
    for celda in ws[3]:
        celda.font = Font(bold=True)
        celda.fill = gris
    verde = PatternFill('solid', fgColor='E2EFDA')
    for f in filas:
        ws.append([f['trabajador'], f['dni'], f['epp'], len(f['entregas'])] +
                  [texto_de_entrega(e) for e in f['entregas']])
        ultima = ws.cell(row=ws.max_row, column=4 + len(f['entregas']))
        ultima.fill = verde
    anchos = [30, 12, 28, 10] + [30] * columnas
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    for fila in ws.iter_rows(min_row=4):
        for celda in fila:
            celda.alignment = Alignment(vertical='top')
    ws.freeze_panes = 'E4'
    ws.append([])
    ws.append(['Días: lo que duró cada entrega hasta la siguiente del mismo EPP. '
               'En verde, la que sigue en uso (cuenta hasta hoy).'])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
