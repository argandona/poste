"""Camiones: kilometraje diario, cargas de combustible e informe mensual.

Pedido del usuario el 2026-10-10:
- Solo el **chofer** (rol nuevo) registra: es el único que conduce. El chofer
  no es fijo: cada salida elige el camión.
- Al **salir** anota el km con foto del tablero y a qué va: SST del
  responsable del camión, trabajo interno de la empresa, o ambos. Al
  **guardar** el camión cierra con km y foto.
- Cada **carga de combustible** lleva km, foto del tablero, cantidad y monto.
  Los camiones usan gasolina, diésel, GLP o GNV.
- **Informe mensual** por camión para el SuperAdmin y el Coordinador, con
  alerta cuando el rendimiento cae más de 15 % bajo su promedio.

El rendimiento: entre dos cargas de tanque lleno es exacto; si en el mes no
hay dos, se aproxima con km del mes / combustible cargado en el mes.
"""
import io
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation

from django.db import models, transaction
from django.utils import timezone
from rest_framework import permissions, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied
from rest_framework.response import Response

from .errores import ErrorNegocio
from .fotos import guardar_foto, leer_base64, responder_foto
from .models import (
    SST, Camion, CargaCombustible, JornadaCamion, JornadaSST, Rol, SSTEncargado,
    Usuario, UsuarioCamion,
)

# Por debajo de este tanto de su promedio, el rendimiento avisa.
UMBRAL_ALERTA = Decimal('0.85')
# Cuántos tramos de tanque lleno anteriores hacen el promedio de la alerta.
TRAMOS_PROMEDIO = 5
# Cuántos meses anteriores hacen el promedio del informe.
MESES_PROMEDIO = 3


# ── Ayudas ───────────────────────────────────────────────────────────────────

def _actor(request):
    return Usuario.objects.filter(pk=request.user.id_usuario).first()


def _km(valor, que='el km'):
    try:
        km = int(str(valor).strip())
    except (TypeError, ValueError):
        raise ErrorNegocio(f'Escribe {que} del tablero.')
    if km < 0:
        raise ErrorNegocio(f'{que.capitalize()} no puede ser negativo.')
    return km


def _decimal(valor, que, decimales):
    try:
        d = Decimal(str(valor).strip().replace(',', '.'))
    except (InvalidOperation, AttributeError):
        d = None
    if d is None or not d.is_finite() or d <= 0:
        raise ErrorNegocio(f'Escribe {que}.')
    return d.quantize(Decimal(1).scaleb(-decimales))


def _foto_obligatoria(valor):
    foto = leer_base64(valor)
    if not foto:
        raise ErrorNegocio('Toma la foto del tablero.')
    return foto


def ultimo_km(camion):
    """El km más alto que se le conoce al camión: salidas, cierres y cargas."""
    j = JornadaCamion.objects.filter(camion=camion).aggregate(
        s=models.Max('km_salida'), c=models.Max('km_cierre'))
    c = CargaCombustible.objects.filter(camion=camion).aggregate(k=models.Max('km'))
    return max([v for v in (j['s'], j['c'], c['k']) if v is not None], default=None)


def responsable_de(camion, fecha):
    asignacion = (UsuarioCamion.objects
                  .filter(camion=camion, activo=True, fecha_inicio__lte=fecha)
                  .filter(models.Q(fecha_fin__isnull=True) | models.Q(fecha_fin__gte=fecha))
                  .select_related('usuario').first())
    return asignacion.usuario if asignacion else None


def ssts_del_responsable(responsable, fecha):
    """Las SST que el responsable tiene por ejecutar (o que ejecutó hoy)."""
    if responsable is None:
        return SST.objects.none()
    ids = SSTEncargado.objects.filter(usuario=responsable).values_list('sst_id', flat=True)
    return (SST.objects.filter(pk__in=ids)
            .filter(models.Q(fecha_ejecucion__isnull=True) | models.Q(fecha_ejecucion=fecha))
            .select_related('actividad').order_by('codigo', 'sst'))


def _sst_texto(sst):
    return sst.codigo or sst.sst or f'SST #{sst.pk}'


def datos_jornada(j):
    cargas = list(j.cargas.all())
    return {
        'id_jornada': j.pk,
        'camion': j.camion_id,
        'camion_placa': j.camion.placa,
        'combustible': j.camion.combustible,
        'unidad': Camion.UNIDADES.get(j.camion.combustible, ''),
        'chofer': j.chofer_id,
        'chofer_nombre': j.chofer.nombre,
        'responsable_nombre': j.responsable.nombre if j.responsable_id else '',
        'fecha': j.fecha,
        'km_salida': j.km_salida,
        'hora_salida': j.hora_salida,
        'km_cierre': j.km_cierre,
        'hora_cierre': j.hora_cierre,
        'km_recorridos': km_de_jornada(j, cargas),
        'abierta': j.abierta,
        'trabajo_interno': j.trabajo_interno,
        'observacion': j.observacion,
        'ssts': [{'id_sst': x.sst_id, 'sst': _sst_texto(x.sst),
                  'actividad': x.sst.actividad.nombre if x.sst.actividad_id else ''}
                 for x in j.ssts.all()],
        'cargas': [datos_carga(c) for c in cargas],
    }


def datos_carga(c):
    return {
        'id_carga': c.pk,
        'camion': c.camion_id,
        'camion_placa': c.camion.placa,
        'chofer_nombre': c.chofer.nombre,
        'jornada': c.jornada_id,
        'fecha_hora': c.fecha_hora,
        'km': c.km,
        'combustible': c.combustible,
        'unidad': Camion.UNIDADES.get(c.combustible, ''),
        'cantidad': c.cantidad,
        'monto': c.monto,
        'precio_unitario': (c.monto / c.cantidad).quantize(Decimal('0.01')),
        'tanque_lleno': c.tanque_lleno,
        'grifo': c.grifo,
    }


def km_de_jornada(j, cargas=None):
    """Lo recorrido en la salida. Si sigue abierta, hasta la última carga."""
    if j.km_cierre is not None:
        return j.km_cierre - j.km_salida
    cargas = cargas if cargas is not None else j.cargas.all()
    return max([c.km for c in cargas] + [j.km_salida]) - j.km_salida


def _jornadas_con_todo(qs):
    return (qs.select_related('camion', 'chofer', 'responsable')
            .prefetch_related('ssts__sst__actividad', 'cargas__camion', 'cargas__chofer'))


def _avisar(empresa_id, titulo, cuerpo, data):
    from .fcm import send_notification
    from .views import con_rol
    tokens = list(
        Usuario.objects.filter(activo=True)
        .filter(con_rol(Rol.SUPERADMIN)
                | (con_rol(Rol.COORDINADOR) & models.Q(empresa_id=empresa_id)))
        .exclude(fcm_token__isnull=True).exclude(fcm_token='')
        .values_list('fcm_token', flat=True))
    send_notification(tokens, title=titulo, body=cuerpo, data=data)


# ── Rendimiento ──────────────────────────────────────────────────────────────

def tramos_llenos(camion, combustible, hasta_km=None):
    """Rendimiento entre cada par de cargas de tanque lleno seguidas, en orden
    de km: [(carga que cierra el tramo, km/unidad)]."""
    cargas = list(CargaCombustible.objects
                  .filter(camion=camion, combustible=combustible)
                  .order_by('km', 'fecha_hora'))
    if hasta_km is not None:
        cargas = [c for c in cargas if c.km <= hasta_km]
    tramos, anterior, suma = [], None, Decimal(0)
    for c in cargas:
        if anterior is not None:
            suma += c.cantidad
        if c.tanque_lleno:
            if anterior is not None and suma > 0 and c.km > anterior.km:
                tramos.append((c, Decimal(c.km - anterior.km) / suma))
            anterior, suma = c, Decimal(0)
    return tramos


def revisar_alerta(carga):
    """Al cerrar un tramo de tanque lleno: si rinde menos del 85 % de los
    tramos anteriores, avisa al SuperAdmin y al Coordinador."""
    if not carga.tanque_lleno:
        return None
    tramos = tramos_llenos(carga.camion, carga.combustible, hasta_km=carga.km)
    if not tramos or tramos[-1][0].pk != carga.pk:
        return None
    actual = tramos[-1][1]
    previos = [r for _, r in tramos[:-1]][-TRAMOS_PROMEDIO:]
    if len(previos) < 2:
        return None
    promedio = sum(previos) / len(previos)
    if actual >= promedio * UMBRAL_ALERTA:
        return None
    unidad = Camion.UNIDADES.get(carga.combustible, '')
    caida = (1 - actual / promedio) * 100
    alerta = {
        'rendimiento': round(actual, 2),
        'promedio': round(promedio, 2),
        'caida_pct': round(caida, 1),
        'unidad': unidad,
    }
    _avisar(carga.camion.empresa_id,
            f'Rendimiento bajo: {carga.camion.placa}',
            f'{actual:.1f} km/{unidad} contra {promedio:.1f} de promedio '
            f'({caida:.0f} % menos). Revisa fugas, fallas o la carga.',
            {'tipo': 'rendimiento_bajo', 'camion': str(carga.camion_id),
             'carga': str(carga.pk)})
    return alerta


def resumen_periodo(camion, desde, hasta, hoy=None):
    """Km, combustible y rendimiento de un camión entre dos fechas."""
    hoy = hoy or timezone.localdate()
    todas = list(_jornadas_con_todo(
        JornadaCamion.objects.filter(camion=camion, fecha__lte=hasta))
        .order_by('hora_salida'))
    jornadas, km_jornadas, km_sin_registrar = [], 0, 0
    km_sst, km_interno, sin_cierre = 0, 0, 0
    ssts, choferes, responsables, dias = set(), set(), set(), set()
    anterior = None
    for j in todas:
        dentro = j.fecha >= desde
        if dentro:
            cargas = list(j.cargas.all())
            km = km_de_jornada(j, cargas)
            hueco = 0
            if anterior is not None:
                fin = anterior.km_cierre
                if fin is None:
                    fin = anterior.km_salida + km_de_jornada(anterior)
                hueco = max(0, j.km_salida - fin)
            km_jornadas += km
            km_sin_registrar += hueco
            sst_de_j = {_sst_texto(x.sst) for x in j.ssts.all()}
            if sst_de_j:
                km_sst += km
            else:
                km_interno += km
            ssts |= sst_de_j
            choferes.add(j.chofer.nombre)
            if j.responsable_id:
                responsables.add(j.responsable.nombre)
            dias.add(j.fecha)
            if j.abierta and j.fecha < hoy:
                sin_cierre += 1
            jornadas.append({**datos_jornada(j), 'km_sin_registrar': hueco})
        anterior = j

    cargas = list(CargaCombustible.objects
                  .filter(camion=camion, fecha_hora__date__gte=desde,
                          fecha_hora__date__lte=hasta)
                  .select_related('camion', 'chofer').order_by('km', 'fecha_hora'))
    cantidad = sum((c.cantidad for c in cargas), Decimal(0))
    monto = sum((c.monto for c in cargas), Decimal(0))
    km_total = km_jornadas + km_sin_registrar

    rendimiento, metodo = None, ''
    llenos = [c for c in cargas if c.tanque_lleno]
    if len(llenos) >= 2 and llenos[-1].km > llenos[0].km:
        suma = sum((c.cantidad for c in cargas
                    if llenos[0].km < c.km <= llenos[-1].km), Decimal(0))
        if suma > 0:
            rendimiento = Decimal(llenos[-1].km - llenos[0].km) / suma
            metodo = 'tanque lleno'
    if rendimiento is None and cantidad > 0 and km_total > 0:
        rendimiento = Decimal(km_total) / cantidad
        metodo = 'aproximado'

    return {
        'jornadas': jornadas,
        'cargas': [datos_carga(c) for c in cargas],
        'dias': len(dias),
        'salidas': len(jornadas),
        'km_jornadas': km_jornadas,
        'km_sin_registrar': km_sin_registrar,
        'km_total': km_total,
        'km_sst': km_sst,
        'km_interno': km_interno,
        'ssts_atendidas': sorted(ssts),
        'km_por_sst': round(Decimal(km_sst) / len(ssts), 1) if ssts else None,
        'sin_cierre': sin_cierre,
        'choferes': sorted(choferes),
        'responsables': sorted(responsables),
        'cantidad': cantidad,
        'monto': monto,
        'rendimiento': round(rendimiento, 2) if rendimiento is not None else None,
        'metodo': metodo,
        'costo_km': round(monto / km_total, 2) if km_total and monto else None,
    }


def _mes_anterior(anio, mes, atras):
    total = anio * 12 + (mes - 1) - atras
    return total // 12, total % 12 + 1


def _fin_de_mes(anio, mes):
    a, m = _mes_anterior(anio, mes, -1)
    return date(a, m, 1) - timedelta(days=1)


def informe_mensual(anio, mes, empresa_id=None, camion_id=None, hoy=None):
    desde, hasta = date(anio, mes, 1), _fin_de_mes(anio, mes)
    a0, m0 = _mes_anterior(anio, mes, MESES_PROMEDIO)
    previo_desde, previo_hasta = date(a0, m0, 1), desde - timedelta(days=1)

    camiones = Camion.objects.all().order_by('placa')
    if empresa_id:
        camiones = camiones.filter(empresa_id=empresa_id)
    if camion_id:
        camiones = camiones.filter(pk=camion_id)
    filas = []
    for camion in camiones:
        r = resumen_periodo(camion, desde, hasta, hoy=hoy)
        if not r['salidas'] and not r['cargas'] and not camion.activo:
            continue
        previo = resumen_periodo(camion, previo_desde, previo_hasta, hoy=hoy)
        promedio = previo['rendimiento']
        variacion = (round((r['rendimiento'] / promedio - 1) * 100, 1)
                     if r['rendimiento'] is not None and promedio else None)
        filas.append({
            'camion': camion.pk,
            'placa': camion.placa,
            'descripcion': camion.descripcion,
            'combustible': camion.combustible,
            'unidad': Camion.UNIDADES.get(camion.combustible, ''),
            **r,
            'rendimiento_promedio': promedio,
            'variacion_pct': variacion,
            'alerta': bool(variacion is not None
                           and r['rendimiento'] < promedio * UMBRAL_ALERTA),
        })
    return {'anio': anio, 'mes': mes, 'desde': desde, 'hasta': hasta,
            'camiones': filas}


MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio',
         'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']


def informe_excel(informe):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    gris = PatternFill('solid', fgColor='DDDDDD')
    rojo = PatternFill('solid', fgColor='F8CBAD')
    titulo = f"Camiones - {MESES[informe['mes'] - 1]} {informe['anio']}"

    def hoja(ws, encabezados, filas, anchos):
        ws.append([titulo])
        ws['A1'].font = Font(bold=True, size=13)
        ws.append([])
        ws.append(encabezados)
        for celda in ws[3]:
            celda.font = Font(bold=True)
            celda.fill = gris
        for f in filas:
            ws.append(f)
        for i, ancho in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.freeze_panes = 'A4'

    def n(x):
        return float(x) if isinstance(x, Decimal) else x

    wb = Workbook()
    ws = wb.active
    ws.title = 'Resumen'
    hoja(ws, ['Placa', 'Combustible', 'Responsables', 'Choferes', 'Días',
              'Salidas', 'Km del mes', 'Km SST', 'Km internos',
              'Km sin registrar', 'SST atendidas', 'Km por SST', 'Cantidad',
              'Unidad', 'S/ combustible', 'Rendimiento (km/unidad)', 'Método',
              'Promedio 3 meses', 'Variación %', 'S/ por km', 'Sin cierre',
              'Alerta'],
         [[c['placa'], c['combustible'].upper(), ', '.join(c['responsables']),
           ', '.join(c['choferes']), c['dias'], c['salidas'], c['km_total'],
           c['km_sst'], c['km_interno'], c['km_sin_registrar'],
           len(c['ssts_atendidas']), n(c['km_por_sst']), n(c['cantidad']),
           c['unidad'], n(c['monto']), n(c['rendimiento']), c['metodo'],
           n(c['rendimiento_promedio']), n(c['variacion_pct']), n(c['costo_km']),
           c['sin_cierre'], 'SÍ' if c['alerta'] else '']
          for c in informe['camiones']],
         [11, 12, 24, 24, 7, 8, 11, 9, 11, 13, 12, 11, 10, 8, 14, 14, 12, 14, 11, 10, 10, 8])
    for fila, c in zip(ws.iter_rows(min_row=4), informe['camiones']):
        if c['alerta']:
            for celda in fila:
                celda.fill = rojo

    hora = lambda d: timezone.localtime(d).strftime('%H:%M') if d else ''
    hoja(wb.create_sheet('Salidas'),
         ['Placa', 'Fecha', 'Chofer', 'Responsable', 'Salida', 'Km salida',
          'Cierre', 'Km cierre', 'Km recorridos', 'Km sin registrar', 'SST',
          'Trabajo interno', 'Observación'],
         [[c['placa'], j['fecha'], j['chofer_nombre'], j['responsable_nombre'],
           hora(j['hora_salida']), j['km_salida'], hora(j['hora_cierre']),
           j['km_cierre'] if j['km_cierre'] is not None else 'SIN CIERRE',
           j['km_recorridos'], j['km_sin_registrar'],
           ', '.join(s['sst'] for s in j['ssts']), j['trabajo_interno'],
           j['observacion']]
          for c in informe['camiones'] for j in c['jornadas']],
         [11, 11, 22, 22, 8, 10, 8, 11, 12, 14, 24, 30, 30])
    hoja(wb.create_sheet('Combustible'),
         ['Placa', 'Fecha', 'Hora', 'Chofer', 'Km', 'Combustible', 'Cantidad',
          'Unidad', 'S/', 'S/ por unidad', 'Tanque lleno', 'Grifo'],
         [[c['placa'], timezone.localtime(k['fecha_hora']).date(),
           hora(k['fecha_hora']), k['chofer_nombre'], k['km'],
           k['combustible'].upper(), n(k['cantidad']), k['unidad'],
           n(k['monto']), n(k['precio_unitario']),
           'Sí' if k['tanque_lleno'] else 'No', k['grifo']]
          for c in informe['camiones'] for k in c['cargas']],
         [11, 11, 7, 22, 10, 12, 10, 8, 10, 12, 12, 24])

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ── API ──────────────────────────────────────────────────────────────────────

def _camiones_de(actor):
    """Los camiones que puede tomar: los de su empresa; el SuperAdmin, todos."""
    qs = Camion.objects.all()
    return qs if actor.es_superadmin() else qs.filter(empresa_id=actor.empresa_id)


class JornadaCamionViewSet(viewsets.ViewSet):
    """Salidas del camión. El chofer sale, cambia a qué va y cierra; el
    SuperAdmin y el Coordinador miran todo y sacan el informe."""
    permission_classes = [permissions.IsAuthenticated]

    @staticmethod
    def _visibles(actor):
        qs = JornadaCamion.objects.all()
        if actor.puede_ver_informe_camiones():
            if not actor.es_superadmin():
                qs = qs.filter(camion__empresa_id=actor.empresa_id)
            return qs
        if actor.puede_manejar():
            return qs.filter(chofer=actor)
        raise PermissionDenied('Tu rol no registra ni ve el kilometraje.')

    @staticmethod
    def _del_chofer(request):
        actor = _actor(request)
        if not actor.puede_manejar():
            raise PermissionDenied('Solo el chofer registra el camión.')
        return actor

    def list(self, request):
        """GET /api/jornadas-camion/?anio=&mes=&camion=&mias=1"""
        actor = _actor(request)
        qs = self._visibles(actor)
        p = request.query_params
        if p.get('mias'):
            qs = qs.filter(chofer=actor)
        if p.get('anio') and p.get('mes'):
            qs = qs.filter(fecha__year=p['anio'], fecha__month=p['mes'])
        if p.get('camion'):
            qs = qs.filter(camion_id=p['camion'])
        return Response([datos_jornada(j) for j in _jornadas_con_todo(qs)[:200]])

    @action(detail=False, methods=['get'])
    def actual(self, request):
        """GET /api/jornadas-camion/actual/ — la salida abierta del chofer y
        los camiones que puede tomar, con su último km."""
        actor = self._del_chofer(request)
        abierta = (_jornadas_con_todo(JornadaCamion.objects
                                      .filter(chofer=actor, km_cierre__isnull=True))
                   .first())
        en_ruta = {j.camion_id: j.chofer.nombre for j in JornadaCamion.objects
                   .filter(km_cierre__isnull=True).select_related('chofer')}
        hoy = timezone.localdate()
        camiones = []
        for c in _camiones_de(actor).filter(activo=True).order_by('placa'):
            r = responsable_de(c, hoy)
            camiones.append({
                'id_camion': c.pk, 'placa': c.placa, 'descripcion': c.descripcion,
                'combustible': c.combustible,
                'unidad': Camion.UNIDADES.get(c.combustible, ''),
                'ultimo_km': ultimo_km(c),
                'responsable_nombre': r.nombre if r else '',
                'en_ruta_con': en_ruta.get(c.pk, ''),
            })
        ultima = JornadaCamion.objects.filter(chofer=actor).order_by('-hora_salida').first()
        return Response({
            'jornada': datos_jornada(abierta) if abierta else None,
            'camiones': camiones,
            'ultimo_camion': ultima.camion_id if ultima else None,
        })

    @action(detail=False, methods=['get'])
    def ssts(self, request):
        """GET /api/jornadas-camion/ssts/?camion= — las SST del responsable
        del camión, para elegir a qué sale."""
        actor = _actor(request)
        camion = _camiones_de(actor).filter(pk=request.query_params.get('camion')).first()
        if camion is None:
            raise ErrorNegocio('Elige el camión.')
        hoy = timezone.localdate()
        r = responsable_de(camion, hoy)
        return Response({
            'responsable_nombre': r.nombre if r else '',
            'ssts': [{'id_sst': s.pk, 'sst': _sst_texto(s), 'distrito': s.distrito,
                      'actividad': s.actividad.nombre if s.actividad_id else ''}
                     for s in ssts_del_responsable(r, hoy)],
        })

    @staticmethod
    def _motivo(empresa_id, datos):
        ids = {int(x) for x in datos.get('ssts') or []}
        ssts = list(SST.objects.filter(pk__in=ids, empresa_id=empresa_id))
        if len(ssts) != len(ids):
            raise ErrorNegocio('Una de las SST no existe.')
        interno = (datos.get('trabajo_interno') or '').strip()
        if not ssts and not interno:
            raise ErrorNegocio('Elige las SST o describe el trabajo interno.')
        return ssts, interno

    def create(self, request):
        """POST /api/jornadas-camion/ — la salida.
        {camion, km, foto (jpeg base64), ssts: [id], trabajo_interno}"""
        actor = self._del_chofer(request)
        d = request.data
        camion = _camiones_de(actor).filter(pk=d.get('camion'), activo=True).first()
        if camion is None:
            raise ErrorNegocio('Elige el camión.')
        km = _km(d.get('km'))
        ssts, interno = self._motivo(camion.empresa_id, d)
        foto = _foto_obligatoria(d.get('foto'))

        propia = JornadaCamion.objects.filter(chofer=actor, km_cierre__isnull=True).first()
        if propia:
            raise ErrorNegocio(f'Primero cierra la salida del {propia.fecha:%d/%m} '
                               f'con {propia.camion.placa}.')
        otra = (JornadaCamion.objects.filter(camion=camion, km_cierre__isnull=True)
                .select_related('chofer').first())
        if otra:
            raise ErrorNegocio(f'{camion.placa} está en ruta con {otra.chofer.nombre}.')
        anterior = ultimo_km(camion)
        if anterior is not None and km < anterior:
            raise ErrorNegocio(f'El km no puede ser menor que el último registrado '
                               f'de {camion.placa}: {anterior}.')

        archivo = guardar_foto(foto, 'camiones')
        hoy = timezone.localdate()
        with transaction.atomic():
            j = JornadaCamion.objects.create(
                camion=camion, chofer=actor, responsable=responsable_de(camion, hoy),
                fecha=hoy, km_salida=km, foto_salida=archivo, trabajo_interno=interno)
            JornadaSST.objects.bulk_create([JornadaSST(jornada=j, sst=s) for s in ssts])
        return Response(datos_jornada(_jornadas_con_todo(JornadaCamion.objects).get(pk=j.pk)),
                        status=201)

    def _abierta_propia(self, request, pk):
        actor = self._del_chofer(request)
        j = JornadaCamion.objects.filter(pk=pk, chofer=actor).select_related('camion').first()
        if j is None:
            raise ErrorNegocio('Esa salida no es tuya.')
        if not j.abierta:
            raise ErrorNegocio('Esa salida ya está cerrada.')
        return actor, j

    @action(detail=True, methods=['post'])
    def motivo(self, request, pk=None):
        """POST /api/jornadas-camion/{id}/motivo/ {ssts, trabajo_interno} —
        mientras está en ruta, cambiar a qué va."""
        actor, j = self._abierta_propia(request, pk)
        ssts, interno = self._motivo(j.camion.empresa_id, request.data)
        with transaction.atomic():
            j.trabajo_interno = interno
            j.save(update_fields=['trabajo_interno'])
            j.ssts.all().delete()
            JornadaSST.objects.bulk_create([JornadaSST(jornada=j, sst=s) for s in ssts])
        return Response(datos_jornada(_jornadas_con_todo(JornadaCamion.objects).get(pk=j.pk)))

    @action(detail=True, methods=['post'])
    def cerrar(self, request, pk=None):
        """POST /api/jornadas-camion/{id}/cerrar/ {km, foto, observacion} — al
        guardar el camión."""
        _, j = self._abierta_propia(request, pk)
        km = _km(request.data.get('km'))
        foto = _foto_obligatoria(request.data.get('foto'))
        minimo = max([j.km_salida] + [c.km for c in j.cargas.all()])
        if km < minimo:
            raise ErrorNegocio(f'El km de cierre no puede ser menor que {minimo}.')
        archivo = guardar_foto(foto, 'camiones')
        j.km_cierre = km
        j.foto_cierre = archivo
        j.hora_cierre = timezone.now()
        j.observacion = (request.data.get('observacion') or '').strip()
        j.save(update_fields=['km_cierre', 'foto_cierre', 'hora_cierre', 'observacion'])
        return Response(datos_jornada(_jornadas_con_todo(JornadaCamion.objects).get(pk=j.pk)))

    @action(detail=True, methods=['get'])
    def foto(self, request, pk=None):
        """GET /api/jornadas-camion/{id}/foto/?cual=salida|cierre"""
        j = self._visibles(_actor(request)).filter(pk=pk).first()
        nombre = None
        if j is not None:
            nombre = j.foto_cierre if request.query_params.get('cual') == 'cierre' else j.foto_salida
        if not nombre:
            return Response({'detail': 'Sin foto.'}, status=404)
        return responder_foto(nombre)

    @action(detail=False, methods=['get'])
    def informe(self, request):
        """GET /api/jornadas-camion/informe/?anio=&mes=[&camion=][&formato=xlsx]"""
        from django.http import HttpResponse
        actor = _actor(request)
        if not actor.puede_ver_informe_camiones():
            raise PermissionDenied('El informe lo ven el SuperAdmin y el Coordinador.')
        hoy = timezone.localdate()
        try:
            anio = int(request.query_params.get('anio') or hoy.year)
            mes = int(request.query_params.get('mes') or hoy.month)
            date(anio, mes, 1)
        except ValueError:
            raise ErrorNegocio('Mes inválido.')
        informe = informe_mensual(
            anio, mes, empresa_id=None if actor.es_superadmin() else actor.empresa_id,
            camion_id=request.query_params.get('camion') or None)
        if request.query_params.get('formato') == 'xlsx':
            r = HttpResponse(
                informe_excel(informe),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            r['Content-Disposition'] = f'attachment; filename="camiones_{anio}_{mes:02d}.xlsx"'
            return r
        return Response(informe)


class CargaCombustibleViewSet(viewsets.ViewSet):
    """Cargas de combustible. Las registra el chofer en ruta; las miran el
    SuperAdmin y el Coordinador."""
    permission_classes = [permissions.IsAuthenticated]

    @staticmethod
    def _visibles(actor):
        qs = CargaCombustible.objects.select_related('camion', 'chofer')
        if actor.puede_ver_informe_camiones():
            if not actor.es_superadmin():
                qs = qs.filter(camion__empresa_id=actor.empresa_id)
            return qs
        if actor.puede_manejar():
            return qs.filter(chofer=actor)
        raise PermissionDenied('Tu rol no registra combustible.')

    def list(self, request):
        """GET /api/cargas-combustible/?anio=&mes=&camion="""
        qs = self._visibles(_actor(request))
        p = request.query_params
        if p.get('anio') and p.get('mes'):
            qs = qs.filter(fecha_hora__year=p['anio'], fecha_hora__month=p['mes'])
        if p.get('camion'):
            qs = qs.filter(camion_id=p['camion'])
        return Response([datos_carga(c) for c in qs[:200]])

    def create(self, request):
        """POST /api/cargas-combustible/
        {km, cantidad, monto, tanque_lleno, grifo, foto, combustible?}
        Va sobre la salida abierta del chofer. Si el camión aún no tiene
        combustible definido, el chofer lo dice y queda puesto."""
        actor = _actor(request)
        if not actor.puede_manejar():
            raise PermissionDenied('Solo el chofer registra el combustible.')
        j = (JornadaCamion.objects.filter(chofer=actor, km_cierre__isnull=True)
             .select_related('camion').first())
        if j is None:
            raise ErrorNegocio('Primero registra la salida del camión.')
        camion = j.camion
        d = request.data
        combustible = camion.combustible or (d.get('combustible') or '')
        if combustible not in Camion.UNIDADES:
            raise ErrorNegocio(f'Elige el combustible de {camion.placa}.')
        km = _km(d.get('km'))
        anterior = ultimo_km(camion)
        if anterior is not None and km < anterior:
            raise ErrorNegocio(f'El km no puede ser menor que el último registrado: {anterior}.')
        unidad = Camion.UNIDADES[combustible]
        cantidad = _decimal(d.get('cantidad'), f'la cantidad en {unidad}', 3)
        monto = _decimal(d.get('monto'), 'el monto en S/', 2)
        foto = _foto_obligatoria(d.get('foto'))
        archivo = guardar_foto(foto, 'combustible')
        with transaction.atomic():
            if not camion.combustible:
                camion.combustible = combustible
                camion.save(update_fields=['combustible'])
            carga = CargaCombustible.objects.create(
                camion=camion, chofer=actor, jornada=j, km=km,
                combustible=combustible, cantidad=cantidad, monto=monto,
                tanque_lleno=bool(d.get('tanque_lleno')),
                grifo=(d.get('grifo') or '').strip()[:150], foto=archivo)
        alerta = revisar_alerta(carga)
        return Response({**datos_carga(carga), 'alerta': alerta}, status=201)

    @action(detail=True, methods=['get'])
    def foto(self, request, pk=None):
        """GET /api/cargas-combustible/{id}/foto/"""
        c = self._visibles(_actor(request)).filter(pk=pk).first()
        if c is None or not c.foto:
            return Response({'detail': 'Sin foto.'}, status=404)
        return responder_foto(c.foto)
