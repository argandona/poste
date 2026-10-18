"""Matriz del saldo de un camión: de dónde sale cada número.

Pedido del usuario el 2026-10-16: un Excel por camión con una fila por
material y, de izquierda a derecha y en orden de fecha, cada movimiento que
tocó el saldo (lo que entró por pedido o asignación, lo que salió por
liquidación o devolución) y al final el saldo que debería tener.

Lo que mueve el saldo del camión, y de dónde se lee:
- **Pedido aprobado** (+): `DetallePedido.cantidad_aprobada`, a la fecha de
  aprobación.
- **Asignación de agregado** (+): `AsignacionAgregado` no anulada (la anulada
  devolvió las bolsas, así que neta es cero). Las del mismo día van juntas en
  una columna.
- **Liquidación** (−): `ConsumoMaterialSuministro`. No guarda el camión: se
  descontó del camión que tenía asignado quien liquidó ese día, así que se
  toman los consumos de los responsables del camión dentro de su periodo.
- **Consumo por Excel aprobado** (−): `DetalleConsumo` (la carga antigua).
- **Devolución aprobada** (−): `DetalleDevolucion.cantidad_aprobada`.
- **Inventario cerrado**: deja el saldo en lo contado. Cada inventario cierra
  una pestaña (pedido del mismo día): saldo inicial = lo contado en el
  anterior, los movimientos, el saldo calculado, lo contado y la diferencia,
  que es el faltante o sobrante de ese inventario. La última pestaña va hasta
  hoy y cierra con el saldo del sistema. Si pasan más de 2 meses sin
  inventario, se parte por mes para que no crezca sin fin.
- **Traspaso**: no mueve material, cambia el responsable. Va como columna
  vacía para que se vea desde cuándo responde cada uno.

Lo del mismo día del inventario cuenta antes del conteo: los pedidos y las
devoluciones solo tienen fecha, no hora.
"""
import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal

from django.db.models import Q
from django.utils import timezone

from .models import (
    AsignacionAgregado, ConsumoMaterialSuministro, DetalleConsumo,
    DetalleDevolucion, DetalleInventario, DetallePedido, Inventario, Material,
    StockCamion,
    TraspasoCamion, UsuarioCamion,
)

ENTRA, SALE, INFO = 'entra', 'sale', 'info'


@dataclass
class Movimiento:
    fecha: date
    tipo: str            # "Pedido", "Liquidación", ...
    documento: str       # "N° 123", "SST 1234 · poste 567"
    responsable: str
    clase: str           # ENTRA, SALE o INFO
    cantidades: dict = field(default_factory=dict)  # material_id -> Decimal (con signo)
    orden: int = 0


def _dia(valor):
    if isinstance(valor, datetime):
        return timezone.localdate(valor) if timezone.is_aware(valor) else valor.date()
    return valor


def _periodos(camion):
    """Quién fue responsable del camión y entre qué fechas."""
    return list(UsuarioCamion.objects.filter(camion=camion)
                .select_related('usuario').order_by('fecha_inicio'))


def _sst_de(liq):
    if liq.sst_externo:
        return liq.sst_externo
    if liq.suministro_id:
        for rel in liq.suministro.sst_suministros.all():
            return rel.sst.codigo or rel.sst.sst or ''
    return ''


def movimientos(camion):
    """Todos los movimientos del camión, del más viejo al más nuevo."""
    movs = {}

    def mov(clave, **datos):
        if clave not in movs:
            movs[clave] = Movimiento(**datos)
        return movs[clave]

    def sumar(m, material_id, cantidad):
        m.cantidades[material_id] = m.cantidades.get(material_id, Decimal('0')) + cantidad

    for d in (DetallePedido.objects
              .filter(pedido__camion=camion, pedido__estado='aprobado',
                      cantidad_aprobada__gt=0)
              .select_related('pedido__usuario')):
        p = d.pedido
        m = mov(('pedido', p.pk), fecha=p.fecha_aprobacion or p.fecha, tipo='Pedido',
                documento=f'N° {p.pk}', responsable=p.usuario.nombre, clase=ENTRA,
                orden=1)
        sumar(m, d.material_id, d.cantidad_aprobada)

    # Las asignaciones van una por material; en la matriz se juntan las del
    # mismo día en una sola columna, con todos sus números.
    for a in (AsignacionAgregado.objects
              .filter(camion=camion, anulada=False, cantidad__gt=0)
              .select_related('usuario').order_by('fecha', 'pk')):
        dia = _dia(a.fecha)
        m = mov(('agregado', dia), fecha=dia, tipo='Asignación', documento='',
                responsable='', clase=ENTRA, orden=2)
        m.documento = f'{m.documento}, {a.pk}' if m.documento else f'N° {a.pk}'
        if a.usuario.nombre not in m.responsable.split(', '):
            m.responsable = ', '.join(x for x in (m.responsable, a.usuario.nombre) if x)
        sumar(m, a.material_id, a.cantidad)

    filtro = Q(pk__in=[])
    for p in _periodos(camion):
        rango = Q(usuario_id=p.usuario_id, fecha__gte=p.fecha_inicio)
        if p.fecha_fin:
            rango &= Q(fecha__lte=p.fecha_fin)
        filtro |= rango
    for c in (ConsumoMaterialSuministro.objects.filter(filtro, cantidad__gt=0)
              .select_related('liquidacion__suministro', 'liquidacion__tipo_trabajo',
                              'usuario')
              .prefetch_related('liquidacion__suministro__sst_suministros__sst')):
        liq = c.liquidacion
        poste = (liq.suministro.numero_suministro if liq.suministro_id
                 else liq.suministro_externo)
        sst = _sst_de(liq)
        documento = ' · '.join(x for x in (f'SST {sst}' if sst else '',
                                           f'poste {poste}' if poste else '') if x)
        m = mov(('liquidacion', liq.pk), fecha=c.fecha, tipo='Liquidación',
                documento=documento or f'N° {liq.pk}', responsable=c.usuario.nombre,
                clase=SALE, orden=3)
        sumar(m, c.material_id, -c.cantidad)

    for d in (DetalleConsumo.objects
              .filter(consumo__camion=camion, consumo__upload__estado='aprobado')
              .select_related('consumo__usuario_consume')):
        co = d.consumo
        m = mov(('consumo', co.pk), fecha=co.fecha, tipo='Consumo (Excel)',
                documento=f'poste {co.suministro}' if co.suministro else f'N° {co.pk}',
                responsable=co.usuario_consume.nombre, clase=SALE, orden=4)
        sumar(m, d.material_id, -Decimal(d.cantidad))

    for d in (DetalleDevolucion.objects
              .filter(devolucion__camion=camion, devolucion__estado='aprobado',
                      cantidad_aprobada__gt=0)
              .select_related('devolucion__usuario')):
        dv = d.devolucion
        m = mov(('devolucion', dv.pk), fecha=dv.fecha_aprobacion or dv.fecha,
                tipo='Devolución', documento=f'N° {dv.pk}',
                responsable=dv.usuario.nombre, clase=SALE, orden=5)
        sumar(m, d.material_id, -d.cantidad_aprobada)

    for t in (TraspasoCamion.objects.filter(camion=camion)
              .select_related('usuario_entrega', 'usuario_recibe')):
        mov(('traspaso', t.pk), fecha=t.fecha, tipo='Traspaso',
            documento=f'{t.usuario_entrega.nombre} → {t.usuario_recibe.nombre}',
            responsable=t.usuario_recibe.nombre, clase=INFO, orden=7)

    return sorted(movs.values(), key=lambda m: (m.fecha, m.orden, m.documento))


@dataclass
class Periodo:
    """Una pestaña: los movimientos entre dos cortes y cómo termina."""
    nombre: str
    movs: list
    inicial: dict           # material_id -> saldo con que arranca
    cierre: str             # INVENTARIO, CORTE o ACTUAL
    contado: dict = field(default_factory=dict)  # lo contado, si cierra con inventario
    inventario: object = None
    hasta: date = None


INVENTARIO, CORTE, ACTUAL = 'inventario', 'corte', 'actual'
# Si pasan más meses que esto sin inventario, el periodo se parte por mes.
MESES_SIN_INVENTARIO = 2
MESES_CORTOS = ['ene', 'feb', 'mar', 'abr', 'may', 'jun', 'jul', 'ago', 'sep',
                'oct', 'nov', 'dic']


def _inventarios(camion):
    """Los inventarios cerrados del camión, con lo contado por material."""
    contados = defaultdict(dict)
    for d in DetalleInventario.objects.filter(inventario__camion=camion,
                                              inventario__estado='cerrado'):
        contados[d.inventario_id][d.material_id] = d.cantidad_fisica
    return [(inv, contados[inv.pk]) for inv in
            Inventario.objects.filter(camion=camion, estado='cerrado')
            .select_related('usuario').order_by('fecha', 'pk')]


def _mes(d):
    return d.year * 12 + d.month - 1


def _por_mes(movs, desde, hasta):
    """Parte [desde, hasta] por mes si pasa de MESES_SIN_INVENTARIO meses.
    Devuelve [(movimientos, último día del tramo)]."""
    if desde is None or _mes(hasta) - _mes(desde) + 1 <= MESES_SIN_INVENTARIO:
        return [(movs, hasta)]
    tramos = []
    for k in range(_mes(desde), _mes(hasta) + 1):
        anio, mes0 = divmod(k, 12)
        dentro = [m for m in movs if _mes(m.fecha) == k]
        if k < _mes(hasta):
            siguiente = date(anio + (mes0 == 11), (mes0 + 1) % 12 + 1, 1)
            ultimo = siguiente - timedelta(days=1)
        else:
            ultimo = hasta
        tramos.append((dentro, ultimo))
    return tramos


def periodos(camion):
    """La historia del camión partida en pestañas: una por inventario, y la
    última hasta hoy. Cada una arranca con lo contado en el inventario anterior
    (los materiales que no se contaron siguen con su saldo calculado)."""
    hoy = timezone.localdate()
    saldo = defaultdict(lambda: Decimal('0'))
    resultado = []
    previo = None   # el inventario que abre el periodo

    def agregar(tramo_movs, hasta, cierre, inv=None, contado=None):
        inicial = {k: v for k, v in saldo.items() if v}
        for m in tramo_movs:
            for mid, c in m.cantidades.items():
                saldo[mid] += c
        if cierre == INVENTARIO:
            nombre = (f'Hasta Inv {inv.pk} ({inv.fecha:%d-%m-%y})' if previo is None
                      else f'Inv {previo.pk} a Inv {inv.pk} ({inv.fecha:%d-%m-%y})')
        elif cierre == CORTE:
            nombre = f'{MESES_CORTOS[hasta.month - 1]} {hasta.year} sin inventario'
        else:
            nombre = f'Actual (desde Inv {previo.pk})' if previo is not None else 'Actual'
        resultado.append(Periodo(nombre=nombre, movs=tramo_movs, inicial=inicial,
                                 cierre=cierre, contado=contado or {},
                                 inventario=inv, hasta=hasta))
        if cierre == INVENTARIO:
            # Desde aquí manda lo contado.
            for mid, cantidad in (contado or {}).items():
                saldo[mid] = cantidad

    def tramos_hasta(lista, hasta, cierre_final, inv=None, contado=None):
        desde = lista[0].fecha if lista else (previo.fecha if previo is not None else None)
        partes = _por_mes(lista, desde, hasta)
        for i, (parte, fin) in enumerate(partes):
            if i == len(partes) - 1:
                agregar(parte, fin, cierre_final, inv, contado)
            else:
                agregar(parte, fin, CORTE)

    restantes = movimientos(camion)
    for inv, contado in _inventarios(camion):
        # Lo del mismo día del inventario cuenta antes del conteo.
        antes = [m for m in restantes if m.fecha <= inv.fecha]
        restantes = [m for m in restantes if m.fecha > inv.fecha]
        tramos_hasta(antes, inv.fecha, INVENTARIO, inv, contado)
        previo = inv
    tramos_hasta(restantes, hoy, ACTUAL)
    return resultado


def _filas(periodo, actual, materiales):
    """Las filas de una pestaña: por material, saldo inicial, cada movimiento,
    totales y cómo cierra (lo contado o el saldo del sistema)."""
    ids = set(periodo.inicial) | set(periodo.contado)
    for m in periodo.movs:
        ids |= set(m.cantidades)
    if periodo.cierre == ACTUAL:
        ids |= {k for k, v in actual.items() if v}
    filas = []
    for mid in sorted(ids, key=lambda i: materiales[i].matricula):
        inicial = periodo.inicial.get(mid, Decimal('0'))
        celdas = [m.cantidades.get(mid) for m in periodo.movs]
        entra = sum((c for m, c in zip(periodo.movs, celdas) if c and m.clase == ENTRA),
                    Decimal('0'))
        sale = -sum((c for m, c in zip(periodo.movs, celdas) if c and m.clase == SALE),
                    Decimal('0'))
        calculado = inicial + entra - sale
        if periodo.cierre == INVENTARIO:
            # Lo que no se contó siguió con su saldo calculado.
            cierre = periodo.contado.get(mid, calculado)
        elif periodo.cierre == ACTUAL:
            cierre = actual.get(mid, Decimal('0'))
        else:
            cierre = calculado
        if not (inicial or any(celdas) or calculado or cierre):
            continue
        filas.append({
            'matricula': materiales[mid].matricula,
            'descripcion': materiales[mid].descripcion,
            'inicial': inicial, 'celdas': celdas, 'entra': entra, 'sale': sale,
            'calculado': calculado, 'cierre': cierre,
            'diferencia': cierre - calculado,
        })
    return filas


def matriz(camion):
    """[(periodo, filas)] para cada pestaña, de la más vieja a la actual."""
    lista = periodos(camion)
    actual = {s.material_id: s.cantidad
              for s in StockCamion.objects.filter(camion=camion)}
    ids = set(actual)
    for p in lista:
        ids |= set(p.inicial) | set(p.contado)
        for m in p.movs:
            ids |= set(m.cantidades)
    materiales = Material.objects.in_bulk(ids)
    return [(p, _filas(p, actual, materiales)) for p in lista]


def _nombre_hoja(nombre, usados):
    """Excel: 31 caracteres, sin / \\ ? * [ ] : y sin repetir."""
    limpio = ''.join('-' if ch in '/\\?*[]:' else ch for ch in nombre)[:31]
    base, k = limpio, 2
    while limpio in usados:
        sufijo = f' ({k})'
        limpio = base[:31 - len(sufijo)] + sufijo
        k += 1
    usados.add(limpio)
    return limpio


def matriz_excel(camion):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    pestanas = matriz(camion)
    colores = {ENTRA: 'C6EFCE', SALE: 'F8CBAD', INFO: 'DDDDDD'}
    gris = PatternFill('solid', fgColor='DDDDDD')
    azul = PatternFill('solid', fgColor='BDD7EE')
    amarillo = PatternFill('solid', fgColor='FFEB9C')
    rojo = PatternFill('solid', fgColor='FF9999')
    negrita = Font(bold=True)
    centro = Alignment(horizontal='center', vertical='center', wrap_text=True)
    fina = Side(style='thin', color='999999')
    borde = Border(left=fina, right=fina, top=fina, bottom=fina)

    def n(x):
        return float(x) if isinstance(x, Decimal) else x

    def encabezado(celda, relleno):
        celda.fill = relleno
        celda.font = negrita
        celda.alignment = centro
        celda.border = borde

    responsable = (UsuarioCamion.objects.filter(camion=camion, activo=True)
                   .filter(Q(fecha_fin__isnull=True) | Q(fecha_fin__gte=date.today()))
                   .select_related('usuario').first())
    quien = responsable.usuario.nombre if responsable else 'sin asignar'

    wb = Workbook()
    wb.remove(wb.active)
    usados = set()
    for periodo, filas in pestanas:
        ws = wb.create_sheet(_nombre_hoja(periodo.nombre, usados))
        ws['A1'] = f'Saldo del camión {camion.placa} · {periodo.nombre}'
        ws['A1'].font = Font(bold=True, size=13)
        if periodo.cierre == INVENTARIO:
            inv = periodo.inventario
            explica = (f'Cierra con el inventario N° {inv.pk} del {inv.fecha:%d/%m/%Y} '
                       f'({inv.usuario.nombre}). Diferencia = contado − calculado.')
            cierre = ['Contado en inventario', 'Diferencia (faltante −)']
        elif periodo.cierre == ACTUAL:
            explica = (f'Responsable: {quien} · Al {date.today():%d/%m/%Y}. '
                       f'Diferencia = saldo en sistema − calculado.')
            cierre = ['Saldo en sistema', 'Diferencia']
        else:
            explica = ('Corte de mes: pasaron más de 2 meses sin inventario. El saldo '
                       'calculado pasa a la pestaña siguiente.')
            cierre = []
        ws['A2'] = explica + ' Verde entra, rojo sale, gris traspaso.'

        for fila, (a, b) in zip((4, 5, 6), (('', ''), ('', ''), ('Matrícula', 'Material'))):
            for col, valor in ((1, a), (2, b)):
                celda = ws.cell(fila, col, valor)
                celda.fill = gris
                celda.font = negrita
        ws.merge_cells(start_row=4, start_column=3, end_row=6, end_column=3)
        encabezado(ws.cell(4, 3, 'Saldo inicial'), amarillo)
        ws.column_dimensions['C'].width = 11
        inicio = 4
        for i, m in enumerate(periodo.movs):
            col = inicio + i
            relleno = PatternFill('solid', fgColor=colores[m.clase])
            for fila, valor in ((4, m.fecha.strftime('%d/%m/%Y')), (5, m.tipo),
                                (6, f'{m.documento}\n{m.responsable}')):
                encabezado(ws.cell(fila, col, valor), relleno)
            ws.column_dimensions[get_column_letter(col)].width = 14
        fin = inicio + len(periodo.movs)
        titulos = ['Total entró', 'Total salió', 'Saldo calculado'] + cierre
        for j, titulo in enumerate(titulos):
            col = fin + j
            ws.merge_cells(start_row=4, start_column=col, end_row=6, end_column=col)
            encabezado(ws.cell(4, col, titulo), azul)
            ws.column_dimensions[get_column_letter(col)].width = 12
        ws.row_dimensions[6].height = 45
        ws.column_dimensions['A'].width = 12
        ws.column_dimensions['B'].width = 38

        claves = ['entra', 'sale', 'calculado'] + (['cierre', 'diferencia'] if cierre else [])
        for r, f in enumerate(filas, start=7):
            ws.cell(r, 1, f['matricula'])
            ws.cell(r, 2, f['descripcion'])
            ws.cell(r, 3, n(f['inicial'])).border = borde
            for i, c in enumerate(f['celdas']):
                if c:
                    ws.cell(r, inicio + i, n(c)).border = borde
            for j, clave in enumerate(claves):
                celda = ws.cell(r, fin + j, n(f[clave]))
                celda.border = borde
                if clave in ('calculado', 'cierre'):
                    celda.font = negrita
                if clave == 'diferencia' and f['diferencia']:
                    celda.fill = rojo
        if not filas:
            ws.cell(7, 1, 'Sin movimientos de material en este periodo.')
        ws.freeze_panes = ws.cell(7, inicio)
    # Abre en la pestaña actual.
    wb.active = len(wb.sheetnames) - 1

    # Toda la historia en lista, con el saldo que iba quedando: para filtrar
    # por material o por fecha. El inventario deja el saldo en lo contado.
    wk = wb.create_sheet('Movimientos')
    wk.append(['Fecha', 'Movimiento', 'Documento', 'Responsable', 'Matrícula',
               'Material', 'Cantidad', 'Saldo después'])
    for celda in wk[1]:
        celda.font = negrita
        celda.fill = gris
    ids = {mid for p, _ in pestanas for m in p.movs for mid in m.cantidades}
    ids |= {mid for p, _ in pestanas for mid in p.contado}
    materiales = Material.objects.in_bulk(ids)
    saldo = defaultdict(lambda: Decimal('0'))
    for periodo, _ in pestanas:
        for m in periodo.movs:
            if m.clase == INFO:
                wk.append([m.fecha, m.tipo, m.documento, m.responsable])
                continue
            for mid, c in sorted(m.cantidades.items(),
                                 key=lambda x: materiales[x[0]].matricula):
                saldo[mid] += c
                mat = materiales[mid]
                wk.append([m.fecha, m.tipo, m.documento, m.responsable, mat.matricula,
                           mat.descripcion, n(c), n(saldo[mid])])
        if periodo.cierre == INVENTARIO:
            inv = periodo.inventario
            for mid, contado in sorted(periodo.contado.items(),
                                       key=lambda x: materiales[x[0]].matricula):
                ajuste = contado - saldo[mid]
                saldo[mid] = contado
                if ajuste:
                    mat = materiales[mid]
                    wk.append([inv.fecha, 'Inventario', f'N° {inv.pk} (contado − calculado)',
                               inv.usuario.nombre, mat.matricula, mat.descripcion,
                               n(ajuste), n(contado)])
    for i, ancho in enumerate([11, 16, 30, 24, 12, 38, 10, 13], 1):
        wk.column_dimensions[get_column_letter(i)].width = ancho
    for celda in wk['A'][1:]:
        celda.number_format = 'DD/MM/YYYY'
    wk.freeze_panes = 'A2'
    wk.auto_filter.ref = wk.dimensions

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
